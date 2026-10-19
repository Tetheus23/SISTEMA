"""Classifica cada item cobrado pela CEABS (veículo x mês) em 4 status, cruzando com a frota.

Uso: python3 -I auditar.py <pasta_ceabs>   (ex.: entregas/ceabs; rodar depois de estruturar.py)

Status:
  OK          em uso no mês: com cliente, reserva SEOP/Casa Civil, uso interno com abastecimento ou manutenção curta
  CONTESTAR   cobrado depois do pedido de cancelamento/retirada, depois do cancelamento registrado pela própria
              CEABS, em duplicidade, ou taxa acima do valor combinado
  SEM USO     pago sem uso: vendido, à venda, parado sem cliente, sinistro/oficina há mais de 5 meses,
              uso interno sem abastecer há mais de 60 dias (no mês atual = cancelar)
  VERIFICAR   veículo não localizado na base da frota
Grava em estrutura/: base.csv, contestar.csv, veiculos_status.csv, equipamentos.csv, sinistros_longos.csv
"""
import calendar
import csv
import datetime as dt
import os
import re
import sys
import warnings
from collections import defaultdict

import openpyxl

warnings.filterwarnings('ignore')
BASE = sys.argv[1]
EST = os.path.join(BASE, 'estrutura')
HOJE = dt.date(2026, 10, 19)
DIAS_SINISTRO = 150     # "parado há mais de 5 meses"
DIAS_PARADO = 90        # sem cliente e sem abastecer
DIAS_INTERNO = 60       # uso interno sem abastecer
TAXA_RESCISAO = 60.0    # valor combinado por equipamento
# pedido de cancelamento/retirada enviado à CEABS (assuntos conferidos um a um nos e-mails)
ASSUNTO_PEDIDO = re.compile(r'Retirada de Rastreadores|RETIRADA DE RASTREADOR/|CANCELAMENTO URGENTE|CANCELAMENTO DE CONTRATO'
                            r'|Confirmação da Rescisão', re.I)
# chassi digitado errado na CEABS -> chave da frota (conferido pela placa)
CHASSI_FROTA = {'9535H5TB6PR048130': '9535G5TB6PR048130', '9536E823XPR040055': 'SED3D06',
                '9BM958154NB253139': '9BM958154NR253139', '9BWKB45U7KP027623': '9BWKB45U7KPO27623',
                'B4S511584': 'ESC0004', 'HBZN580NAMAH27211': 'HBZN580NEMAH27211',
                'HBZN580NEMAH26774': 'HBZN580NCMAH26774', '4DFCAP15RB134726': '94DFCAP15RB134726'}


def ler(nome):
    return list(csv.DictReader(open(os.path.join(EST, nome), encoding='utf-8')))


def d(s):
    s = (s or '')[:10]
    return dt.date.fromisoformat(s) if re.match(r'\d{4}-\d{2}-\d{2}$', s) else None


def br(x):
    return x.strftime('%d/%m/%Y') if x else ''


def limites(mes):
    a, m = map(int, mes.split('-'))
    return dt.date(a, m, 1), dt.date(a, m, calendar.monthrange(a, m)[1])


def anexo(sha):
    for r in csv.DictReader(open(os.path.join(EST, 'anexos_unicos.tsv'), encoding='utf-8'), delimiter='\t'):
        if r['sha1'].startswith(sha):
            return os.path.join(BASE, r['arquivo'])


def escreve(nome, linhas):
    with open(os.path.join(EST, nome), 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=list(linhas[0].keys()))
        w.writeheader()
        w.writerows(linhas)


# ---------- frota ----------
frota = {}
for r in ler('frota/veiculos_frota.csv'):
    if r['veiculo_key']:
        frota[r['ch']] = r
extra = {'9535G5TB6PR048130': ('SEH4D59', 'VW/9.170 DRC 4X2', 'sem_destino_registrado', '', ''),
         'SED3D06': ('SED3D06', 'VW/17.190 CRM 4x2 ROB', 'ativo_disponivel_locacao', '', ''),
         '9BM958154NR253139': ('RHS9C37', 'M.BENZ/ATEGO 1719 CL', 'vendido', '2025-02-25', 'Luiz Carlos - Transcar'),
         '9BWKB45U7KPO27623': ('BCO7085', 'VW/NOVA SAVEIRO RB MBVS CAB. SIMPLES', 'ativo_base_grillo', '', ''),
         'ESC0004': ('ESC0004', 'MINI ESCAVADERA HIDRÁULICA', 'ativo_disponivel_locacao', '', ''),
         'HBZN580NEMAH27211': ('RET0041', 'RETROESCAVADEIRA CASE', 'sem_destino_registrado', '', ''),
         'HBZN580NCMAH26774': ('RET0039', 'RETROESCAVADEIRA CASE', 'vendido', '2024-11-22', 'Tarick Correa')}
for ch, chave in CHASSI_FROTA.items():
    if ch not in frota and chave in extra:
        p, mod, sit, venda, comp = extra[chave]
        frota[ch] = dict(veiculo_key=chave, placa=p, marca_modelo=mod, situacao_atual=sit, data_venda=venda,
                         vendido_para=comp, cliente_atual='', local_categoria='')
pass_por = defaultdict(list)
for r in ler('frota/passagens.csv'):
    pass_por[r['veiculo_key']].append(dict(r, e=d(r['entrada']), s=d(r['saida'])))
sin_por = defaultdict(list)
for r in ler('frota/sinistros.csv'):
    sin_por[r['veiculo_key']].append(dict(r, ev=d(r['data_evento']) or d(r['data_solicitacao']), lib=d(r['data_liberacao'])))
comb = defaultdict(set)  # placa -> datas aproximadas de abastecimento (último do mês) e meses com abastecimento
comb_mes = defaultdict(set)
for nome in ('frota/diesel_mensal.csv', 'frota/ticketlog_mensal.csv'):
    for r in ler(nome):
        p = (r['placa'] or '').replace('-', '').upper()
        comb_mes[p].add(r['mes'])
        if nome.startswith('frota/diesel'):
            comb[p].add(d(r['ultimo']))
        else:
            comb[p].add(limites(r['mes'])[1] if r['mes'] < r['ultimo'][:7] else d(r['ultimo']))

# relatório de posição da própria CEABS em 17/07/2026 (e-mail "VEÍCULOS COM RASTREADOR - TICKETLOG")
BASES = re.compile(r'Pedregoso|Camboat|Carlos Seidl|Caju', re.I)
geo = {r['placa']: r for r in ler('frota/geoloc_ceabs_20260717.csv')}
GEO_MESES = ('2026-07', '2026-08')


def em_uso_rastreador(ch, mes):
    """Veículo ligado ou fora das bases da Grillo no relatório da CEABS de 17/07/2026."""
    if mes not in GEO_MESES:
        return None
    for p in placas(ch):
        g = geo.get(p)
        if g and ('On' in g['evento'] or not BASES.search(g['local'] or '')):
            return f'rastreador CEABS em 17/07/2026: {g["evento"]} em {g["local"]}'
    return None


# planilha de veículos de serviço (Monitoramento, 22/08/2026): uso interno Caju/Deodoro
servico = {}
for ws in openpyxl.load_workbook(anexo('ae1cca'), data_only=True, read_only=True):
    for x in list(ws.iter_rows(values_only=True))[1:]:
        if x[0]:
            servico[str(x[0]).replace('-', '').strip().upper()] = f'{x[3]} - {x[4]}'

# ---------- CEABS ----------
itens = [r for r in ler('itens_cobrados.csv') if r['mes_uso'] >= '2024-08']  # movimentação da frota carregada a partir de jun/24
reg = {r['chassi']: r for r in ler('veiculos.csv')}
pedidos = defaultdict(list)
for r in ler('pedidos.csv'):
    if ASSUNTO_PEDIDO.search(r['assunto']) and r['chassi']:
        pedidos[r['chassi']].append((d(r['data']), r['assunto'], r['remetente']))
taxa_mes = {}  # chassi -> 1º mês em que a CEABS cobrou taxa de rescisão (contrato encerrado)
for r in itens:
    if r['tipo'] == 'Taxa de Rescisão':
        taxa_mes[r['chassi']] = min(taxa_mes.get(r['chassi'], '9'), r['mes_uso'])


def placas(ch):
    out = {reg.get(ch, {}).get('placa', '')}
    if ch in frota:
        out.add(frota[ch]['placa'])
    return {p for p in out if p}


def pedido_valido(ch):
    """Primeiro pedido à CEABS depois da instalação vigente. Se o veículo voltou a um cliente depois do pedido,
    o pedido foi abandonado (não houve aceite da rescisão) e vale o pedido seguinte, se houver."""
    inst = d(reg.get(ch, {}).get('instalacao'))
    key = frota.get(ch, {}).get('veiculo_key')
    for p in sorted(p for p in pedidos.get(ch, []) if p[0] and (not inst or p[0] >= inst)):
        voltou = any(x['local_cat'] == 'cliente' and x['e'] and x['e'] > p[0] + dt.timedelta(days=15)
                     for x in pass_por.get(key, []))
        if not voltou:
            return p
    return None


def local_no_mes(key, ini, fim):
    """Passagem com mais dias dentro do mês e, à parte, o cliente com mais dias no mês."""
    melhor, cli = None, None
    for p in pass_por.get(key, []):
        e, s = p['e'] or ini, p['s'] or HOJE
        dias = (min(s, fim) - max(e, ini)).days + 1
        if dias <= 0:
            continue
        if not melhor or dias > melhor[1]:
            melhor = (p, dias)
        if p['local_cat'] == 'cliente' and (not cli or dias > cli[1]):
            cli = (p, dias)
    return melhor, cli


def ultimo_cliente(key, ate):
    ps = [p for p in pass_por.get(key, []) if p['local_cat'] == 'cliente' and (p['e'] or ate) <= ate]
    return max(ps, key=lambda p: p['s'] or HOJE) if ps else None


def ultimo_abast(ch, ate):
    ds = [x for p in placas(ch) for x in comb.get(p, ()) if x and x <= ate]
    return max(ds) if ds else None


def classifica_uso(ch, mes):
    ini, fim = limites(mes)
    fim_ref = min(fim, HOJE)
    f = frota.get(ch)
    if not f:
        return 'VERIFICAR', 'Veículo não localizado na base da frota', '', ''
    key = f['veiculo_key']
    venda = d(f.get('data_venda'))
    if venda and venda < ini:
        return 'SEM USO', f'Vendido em {br(venda)} ({f.get("vendido_para") or "comprador não informado"}) e sem pedido de cancelamento', 'VENDIDO', 'vendido'
    (loc, _), cli = local_no_mes(key, ini, fim) if local_no_mes(key, ini, fim)[0] else ((None, 0), None)
    abast_mes = any(mes in comb_mes.get(p, ()) for p in placas(ch))
    # sinistro aberto há mais de 5 meses
    for s in sin_por.get(key, []):
        if s['ev'] and s['ev'] <= fim and (not s['lib'] or s['lib'] >= ini):
            parado = (min(s['lib'] or HOJE, fim_ref) - s['ev']).days
            if parado > DIAS_SINISTRO and not (em_uso_rastreador(ch, mes) or '').count(': Posição / On'):
                return ('SEM USO', f'Sinistro desde {br(s["ev"])} ({s["cliente_base"] or "-"}), {parado // 30} meses parado'
                        + (f'; liberado em {br(s["lib"])}' if s['lib'] else ''), loc['cliente'] if loc else 'SINISTRO', 'sinistro')
    if cli and cli[1] >= 10:
        return 'OK', f'Com cliente: {cli[0]["cliente"]} ({cli[1]} dias no mês)', cli[0]['cliente'], 'cliente'
    nome = loc['cliente'] if loc else (f.get('cliente_atual') or '')
    cat = loc['local_cat'] if loc else (f.get('local_categoria') or '')
    geo_uso = em_uso_rastreador(ch, mes)
    # reserva SEOP / Casa Civil: passou pelo cliente nos 6 meses anteriores ou volta a ele em até 60 dias
    for p in pass_por.get(key, []):
        if p['local_cat'] == 'cliente' and re.search(r'SEOP|CASA CIVIL', p['cliente'] or '', re.I):
            e, s = p['e'] or ini, p['s'] or HOJE
            if s >= ini - dt.timedelta(days=180) and e <= fim + dt.timedelta(days=60):
                return 'OK', f'Reserva {p["cliente"]} (passagem {br(e)} a {br(p["s"]) or "hoje"})', nome or p['cliente'], 'reserva'
    placa_f = f['placa']
    interno = cat == 'base_grillo' or re.search(r'USO GRILLO|LGR|BASE CAJU', nome or '', re.I) or placa_f in servico
    ult = ultimo_abast(ch, fim_ref)
    if interno:
        rot = servico.get(placa_f, nome or 'base Grillo')
        if ult and (fim_ref - ult).days <= DIAS_INTERNO:
            return 'OK', f'Uso interno ({rot}); abastecido em {br(ult)}', rot, 'interno'
        return 'SEM USO', f'Uso interno ({rot}) sem abastecimento há +{DIAS_INTERNO} dias' + (f' (último {br(ult)})' if ult else ''), rot, 'interno'
    if cat in ('disponivel_venda', 'leilao', 'vendido'):
        return 'SEM USO', f'{nome} (desde {br(loc["e"]) if loc else "-"})', nome, cat
    if geo_uso:
        return 'OK', f'{nome or "Sem local registrado"}, mas em uso: {geo_uso}', nome, cat or 'rastreador'
    if abast_mes:
        return 'OK', f'{nome or "Sem cliente"}, mas abasteceu no mês', nome, cat or 'abastecimento'
    if cat == 'oficina_patio':
        dias = (fim_ref - (loc['e'] or ini)).days
        if dias > DIAS_SINISTRO:
            return 'SEM USO', f'Em oficina/pátio desde {br(loc["e"])} ({dias // 30} meses)', nome, cat
        return 'OK', f'Em manutenção desde {br(loc["e"])}', nome, cat
    uc = ultimo_cliente(key, fim)
    parado_desde = (uc['s'] if uc and uc['s'] else None) or (loc['e'] if loc else None)
    if parado_desde:
        dias = (fim_ref - parado_desde).days
        txt = f'{nome or "Sem cliente"}; sem cliente desde {br(parado_desde)}' + (f' (último: {uc["cliente"]})' if uc else '')
        return ('SEM USO' if dias > DIAS_PARADO else 'OK'), txt + f', {dias} dias', nome, cat or 'sem cliente'
    return 'VERIFICAR', f'Sem movimentação registrada no mês ({f.get("situacao_atual")})', nome, cat


# ---------- classificação ----------
base, vistos = [], defaultdict(int)
for r in sorted(itens, key=lambda x: (x['mes_uso'], x['tipo'], x['chassi'])):
    ch, mes, tipo, valor = r['chassi'], r['mes_uso'], r['tipo'], float(r['valor'])
    ini, _ = limites(mes)
    f = frota.get(ch, {})
    placa = r['placa'] or f.get('placa', '')
    ped = pedido_valido(ch)
    canc = d(reg.get(ch, {}).get('cancelamento'))
    prorata = 0 < valor < float(r['mensal'] or 0)
    st, motivo, evid, contestar = 'OK', '', '', 0.0
    vistos[(mes, tipo, ch)] += 1
    uso = classifica_uso(ch, mes) if tipo == 'Principal' else None
    if tipo == 'Principal':
        if vistos[(mes, tipo, ch)] > 1:
            st, motivo, contestar = 'CONTESTAR', 'Cobrado em duplicidade no mesmo mês', valor
        elif ped and ped[0] < ini and not (prorata and (ini - ped[0]).days <= 45) and uso[3] != 'cliente':
            # pró-rata no mês seguinte ao pedido = dias até a CEABS efetivar a rescisão (aceite pelo link): não contesta.
            # Veículo que voltou a cliente depois do pedido: o pedido não foi concluído e o rastreador seguiu em uso.
            st, motivo, contestar = 'CONTESTAR', f'Cancelamento/retirada pedido em {br(ped[0])} e ainda cobrado', valor
            evid = f'E-mail "{ped[1]}" ({br(ped[0])}, {ped[2]})'
        elif canc and canc < ini:
            st, motivo, contestar = 'CONTESTAR', f'CEABS registrou o cancelamento em {br(canc)} e seguiu cobrando', valor
            evid = 'Data de cancelamento na própria planilha de medição'
        elif ch in taxa_mes and taxa_mes[ch] < mes:
            st, motivo, contestar = 'CONTESTAR', f'Taxa de rescisão já cobrada em {taxa_mes[ch]} e mensalidade continua', valor
            evid = f'Aba Taxa de Rescisão da medição {taxa_mes[ch]}'
    elif tipo == 'Taxa de Rescisão' and valor > TAXA_RESCISAO:
        st, motivo, contestar = 'CONTESTAR', f'Taxa de rescisão de R$ {valor:.2f}; o combinado é R$ {TAXA_RESCISAO:.2f} por equipamento', valor - TAXA_RESCISAO
    elif tipo == 'Visita Improdutiva' and vistos[(mes, tipo, ch)] > 1:
        st, motivo, contestar = 'CONTESTAR', 'Visita improdutiva cobrada em duplicidade no mesmo veículo e mês', valor
    if tipo == 'Principal' and st == 'OK':
        st, motivo, local, cat = uso
        if ped and 'sem pedido de cancelamento' in motivo:
            motivo = motivo.replace('sem pedido de cancelamento', f'cancelamento pedido só em {br(ped[0])}')
        if ped and ped[0] < ini and cat == 'cliente':
            motivo += f' (pedido de cancelamento de {br(ped[0])} não concluído; seguiu em uso)'
    else:
        _, m2, local, cat = uso if tipo == 'Principal' else ('', '', '', '')
        if tipo != 'Principal':
            local, cat = 'Taxa / serviço avulso', tipo
            if st == 'OK':
                motivo = f'{tipo} cobrada (cancelamento em {br(d(r["cancelamento"]))}, retirada em {br(d(r["retirada"]))})' \
                    if tipo == 'Taxa de Rescisão' else f'{tipo} cobrada'
        else:
            evid = (evid + ' | ' if evid else '') + f'Situação no mês: {m2}'
    if not evid and ped:
        evid = f'Pedido de cancelamento: {br(ped[0])} ("{ped[1]}")'
    base.append(dict(mes_uso=mes, ref_ceabs=r['ref_ceabs'], tipo=tipo, placa=placa, chassi=ch, codigo_ceabs=r['codigo'],
                     modelo=f.get('marca_modelo', ''), valor=valor, status=st, motivo=motivo, local_no_mes=local,
                     categoria_local=cat, valor_contestar=round(contestar, 2),
                     valor_sem_uso=valor if st == 'SEM USO' else 0.0, pedido_cancelamento=br(ped[0]) if ped else '',
                     cancelamento_ceabs=br(canc), instalacao=br(d(r['instalacao'])), situacao_frota=f.get('situacao_atual', ''),
                     cliente_atual=f.get('cliente_atual', ''), evidencia=evid, planilha_origem=r['arquivo']))
escreve('base.csv', base)

ult_mes = max(b['mes_uso'] for b in base)
escreve('contestar.csv', [b for b in base if b['status'] == 'CONTESTAR'])

# status por veículo no último mês cobrado
porv = {}
for b in base:
    if b['tipo'] == 'Principal' and b['mes_uso'] == ult_mes:
        meses = [x for x in base if x['chassi'] == b['chassi'] and x['tipo'] == 'Principal']
        sem_uso = [x for x in meses if x['status'] == 'SEM USO']
        porv[b['chassi']] = dict(b, meses_cobrados=len(meses), total_pago=round(sum(x['valor'] for x in meses), 2),
                                 meses_sem_uso=len(sem_uso), pago_sem_uso=round(sum(x['valor'] for x in sem_uso), 2))
escreve('veiculos_status.csv', sorted(porv.values(), key=lambda x: (x['status'], x['placa'])))

# sinistros parados há mais de 5 meses (veículos que tiveram rastreador CEABS)
longos = []
for ch, f in frota.items():
    periodos = []
    for s in sorted((x for x in sin_por.get(f['veiculo_key'], []) if x['ev']), key=lambda x: x['ev']):
        if periodos and s['ev'] <= (periodos[-1]['lib'] or HOJE):
            periodos[-1] = dict(periodos[-1], lib=max(periodos[-1]['lib'] or HOJE, s['lib'] or HOJE) if s['lib'] and periodos[-1]['lib'] else None)
        else:
            periodos.append(dict(s))
    for s in periodos:
        if s['ev']:
            dias = ((s['lib'] or HOJE) - s['ev']).days
            if dias > DIAS_SINISTRO:
                cobr = [b for b in base if b['chassi'] == ch and b['tipo'] == 'Principal'
                        and limites(b['mes_uso'])[1] >= s['ev'] and limites(b['mes_uso'])[0] <= (s['lib'] or HOJE)]
                longos.append(dict(placa=f['placa'], chassi=ch, modelo=f['marca_modelo'], cliente=s['cliente_base'],
                                   atendimento=s['atendimento'], data_sinistro=br(s['ev']), entrada_oficina=br(d(s['data_entrada_unidade'])),
                                   liberacao=br(s['lib']) or 'em aberto', dias_parado=dias,
                                   meses_cobrados_no_periodo=len(cobr), valor_cobrado_no_periodo=round(sum(b['valor'] for b in cobr), 2),
                                   ainda_cobrado=('sim' if any(b['mes_uso'] == ult_mes for b in cobr) else 'não')))
escreve('sinistros_longos.csv', sorted(longos, key=lambda x: -x['dias_parado']))
cont = defaultdict(float)
for b in base:
    cont[(b['mes_uso'], b['status'])] += b['valor']
print('último mês:', ult_mes)
for st in ('OK', 'CONTESTAR', 'SEM USO', 'VERIFICAR'):
    n = sum(1 for b in base if b['mes_uso'] == ult_mes and b['status'] == st)
    print(f'  {st:10} {n:4}  R$ {cont[(ult_mes, st)]:9.2f}')
print('contestar total (todos os meses): R$', round(sum(b['valor_contestar'] for b in base), 2),
      '| sem uso total: R$', round(sum(b['valor_sem_uso'] for b in base), 2), '| sinistros longos:', len(longos))

# ---------- cobrança de equipamentos não devolvidos (notificação CEABS de 13/10/2026, fatura a partir de 20/10) ----------
ws = openpyxl.load_workbook(anexo('88c9754b'), data_only=True, read_only=True).active
lin = list(ws.iter_rows(values_only=True))
cab = {c: i for i, c in enumerate(lin[0])}
taxa_paga = {}
for r in itens:
    if r['tipo'] == 'Taxa de Rescisão':
        taxa_paga[r['chassi']] = r['mes_uso']
ret = {}
for x in list(openpyxl.load_workbook(anexo('a5f78e'), data_only=True, read_only=True).active.iter_rows(values_only=True))[1:]:
    if x[2]:
        ret[str(x[2]).upper()] = x  # planilha RETIRADA DE EQUIPAMENTO (Monitoramento)
equip = []
for x in lin[1:]:
    if not x[cab['Valor Faturamento']]:
        continue
    placa_c, chassi_c = x[cab['ds_placa']] or '', x[cab['ds_chassi']] or ''
    ch = chassi_c if chassi_c in reg else next((c for c, v in reg.items() if v['placa'] == placa_c or c == placa_c), chassi_c or placa_c)
    f = frota.get(ch, {})
    venda = d(f.get('data_venda'))
    peds = sorted({p[0] for p in pedidos.get(ch, []) if p[0]})
    tx = taxa_paga.get(ch)
    rr = ret.get(f.get('placa', placa_c).upper())
    situ = rr[6] if rr else ''
    if not venda:
        acao = 'AGENDAR RETIRADA até 20/10 (corporativo@ceabs.com.br): veículo ainda é da Grillo'
        if peds and tx:
            acao += f'; retirada já pedida em {br(peds[0])} e taxa de rescisão cobrada ({tx}): não pagar os R$ 400'
    elif peds and peds[0] < venda and tx:
        acao = (f'CONTESTAR: retirada pedida em {br(peds[0])} e taxa de rescisão já cobrada ({tx}); '
                f'a CEABS não retirou o equipamento antes da venda ({br(venda)})')
    elif peds and peds[0] < venda:
        acao = f'CONTESTAR: retirada pedida em {br(peds[0])}, antes da venda ({br(venda)})'
    else:
        acao = (f'Vendido em {br(venda)}; o equipamento ficou no veículo: combinar a retirada com o comprador '
                f'({f.get("vendido_para") or "-"}) ou aceitar a cobrança')
    equip.append(dict(placa=f.get('placa') or placa_c, chassi=ch, modelo=f.get('marca_modelo', ''),
                      dispositivo=x[cab['ds_modelo_dispositivo']], serie=x[cab['nr_serie_dispositivo']],
                      valor=x[cab['Valor Faturamento']], situacao_frota=f.get('situacao_atual', ''), data_venda=br(venda),
                      comprador=f.get('vendido_para', ''), pedidos_retirada=', '.join(br(p) for p in peds),
                      taxa_rescisao_cobrada=tx or '', retirada_planilha_monitoramento=situ or '', acao=acao))
escreve('equipamentos.csv', equip)
print('equipamentos cobrados:', len(equip), 'R$', sum(e['valor'] for e in equip))
for e in equip:
    print(f"  {e['placa']:9} {e['situacao_frota'][:12]:12} venda {e['data_venda'] or '-':10} pedidos {e['pedidos_retirada'] or '-':24} taxa {e['taxa_rescisao_cobrada'] or '-':8} | {e['acao'][:70]}")

# ---------- retiradas: o que pedir agora e a situação dos já cancelados ----------
mon = {}
for x in list(openpyxl.load_workbook(anexo('3c70b95b'), data_only=True, read_only=True).active.iter_rows(values_only=True))[1:]:
    if x[2]:  # MODELO, Chassi, Placa, LOCAL, DATA/RETIRADA, SITUAÇÃO, CONTATO, PRAZO, DATA/CANCELAMENTO, CONTRATO
        mon[str(x[2]).upper()] = dict(local=x[3] or '', data=x[4], situacao=x[5] or '')
cobrar_equip = {e['chassi'] for e in equip} | {e['placa'] for e in equip}


def placa_de(ch):
    return (frota.get(ch) or {}).get('placa') or reg.get(ch, {}).get('placa') or ch


def data_txt(v):
    # a planilha do Monitoramento traz "04/08/20256" para o RHY1E89 (ano digitado errado)
    return v.strftime('%d/%m/%Y') if hasattr(v, 'strftime') else (str(v).replace('/20256', '/2026') if v else '')


retiradas = []
# A) sem uso no mês atual: pedir cancelamento e retirada
for b in base:
    if b['mes_uso'] == ult_mes and b['tipo'] == 'Principal' and b['status'] == 'SEM USO':
        ch = b['chassi']
        ped = sorted(p for p in pedidos.get(ch, []) if p[0] and p[0] >= limites(ult_mes)[0])
        situ = {'sinistro': 'Sinistrado', 'disponivel_venda': 'À venda', 'vendido': 'Vendido', 'interno': 'Uso interno sem uso'}.get(
            b['categoria_local'], 'Sem cliente')
        acao = (f'Cancelamento já pedido em {br(ped[0][0])}: confirmar o aceite da rescisão e agendar a retirada' if ped
                else 'Solicitar cancelamento e retirada')
        if b['categoria_local'] == 'sinistro':
            acao = 'Confirmar se voltou à operação; se seguir parado, solicitar cancelamento e retirada'
        retiradas.append(dict(grupo='1. Solicitar cancelamento e retirada', placa=b['placa'], chassi=ch, modelo=b['modelo'],
                              situacao=situ, local=b['local_no_mes'], cancelamento=br(ped[0][0]) if ped else '', retirada='', acao=acao))
# B) já cancelados (rescisão cobrada desde jun/25) e pedidos de setembro/26 ainda sem medição
vistos_ret = set()
for r in sorted(itens, key=lambda x: x['mes_uso']):
    if r['tipo'] != 'Taxa de Rescisão' or r['mes_uso'] < '2025-06' or r['chassi'] in vistos_ret:
        continue
    vistos_ret.add(r['chassi'])
    p = placa_de(r['chassi'])
    m = mon.get(p.upper(), {})
    ret_ceabs = d(r['retirada'])
    if ret_ceabs:
        situ, acao = f'Retirado em {br(ret_ceabs)} (medição CEABS)', '-'
    elif 'RETIRADO' in str(m.get('situacao', '')).upper():
        situ, acao = f'Retirado em {data_txt(m.get("data"))} (planilha do Monitoramento)', 'Pedir à CEABS que registre a retirada'
    elif r['chassi'] in cobrar_equip or p in cobrar_equip:
        situ, acao = 'NÃO retirado: está na cobrança de equipamentos da CEABS (R$ 400)', 'Agendar retirada ou tratar a cobrança até 20/10 (aba Equipamentos)'
    else:
        situ, acao = 'Sem confirmação de retirada', 'Informar se o equipamento foi retirado; se não, agendar a retirada'
    if m.get('situacao') and 'RETIRADO' not in str(m['situacao']).upper():
        situ += f' | Monitoramento: {m["situacao"]}'
    retiradas.append(dict(grupo='2. Já cancelados: situação da retirada', placa=p, chassi=r['chassi'],
                          modelo=(frota.get(r['chassi']) or {}).get('marca_modelo', ''), situacao=situ, local=m.get('local', ''),
                          cancelamento=br(d(r['cancelamento'])), retirada=br(ret_ceabs) if ret_ceabs else data_txt(m.get('data')) if 'RETIRADO' in str(m.get('situacao', '')).upper() else '',
                          acao=acao))
ja = {x['chassi'] for x in retiradas}
for e in equip:  # equipamentos que a CEABS vai cobrar e que foram cancelados antes de jun/25
    if e['chassi'] not in ja:
        retiradas.append(dict(grupo='2. Já cancelados: situação da retirada', placa=e['placa'], chassi=e['chassi'], modelo=e['modelo'],
                              situacao='NÃO retirado: está na cobrança de equipamentos da CEABS (R$ 400)', local=e['situacao_frota'],
                              cancelamento=e['pedidos_retirada'].split(',')[0], retirada='',
                              acao='Agendar retirada ou tratar a cobrança até 20/10 (aba Equipamentos)'))
        ja.add(e['chassi'])
for ch, ps in pedidos.items():
    ps = sorted(p for p in ps if p[0] and p[0] >= dt.date(2026, 9, 1))
    if ps and ch not in ja:
        retiradas.append(dict(grupo='3. Cancelamento pedido em set/26', placa=placa_de(ch), chassi=ch,
                              modelo=(frota.get(ch) or {}).get('marca_modelo', ''), situacao=f'Pedido em {br(ps[0][0])} ("{ps[0][1][:40]}")',
                              local=(frota.get(ch) or {}).get('cliente_atual', ''), cancelamento='', retirada='',
                              acao='Confirmar o aceite digital da rescisão (sem ele o contrato segue ativo) e agendar a retirada em até 30 dias'))
        ja.add(ch)
escreve('retiradas.csv', retiradas)
from collections import Counter as _C
print('retiradas:', _C(x['grupo'] for x in retiradas), _C(x['situacao'][:22] for x in retiradas if x['grupo'].startswith('2')))
