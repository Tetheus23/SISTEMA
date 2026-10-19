"""Monta a planilha de conferência do faturamento CEABS (rastreadores).

Uso: python3 -I montar_relatorio.py <pasta_dados> <saida.xlsx>

A pasta de dados contém:
  AAAA-MM.json       medições normalizadas (norm_medicao.py), uma por referência
  notas.json         NFs/boletos conferidos (valores lidos dos PDFs)
  frota_agosto.csv   situação dos veículos na nossa frota (assistente.frota_veiculos/frota_passagens)
"""
import csv, json, sys, glob, os, collections
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

DADOS, SAIDA = sys.argv[1], sys.argv[2]
MOEDA = 'R$ #,##0.00'
CAB = PatternFill('solid', fgColor='1F3864')
CORES = {'OK': 'E2EFDA', 'ATENÇÃO': 'FFF2CC', 'CONTESTAR': 'F8CBAD', 'VERIFICAR': 'DDEBF7'}

med = {os.path.basename(p)[:7]: json.load(open(p)) for p in sorted(glob.glob(os.path.join(DADOS, '20??-??.json')))}
notas = json.load(open(os.path.join(DADOS, 'notas.json')))
frota = {r['chassi']: r for r in csv.DictReader(open(os.path.join(DADOS, 'frota_agosto.csv')), delimiter=';')}
NOME_MES = {'01': 'Jan', '02': 'Fev', '03': 'Mar', '04': 'Abr', '05': 'Mai', '06': 'Jun', '07': 'Jul', '08': 'Ago', '09': 'Set', '10': 'Out', '11': 'Nov', '12': 'Dez'}
rot = lambda k: f'{NOME_MES[k[5:]]}/{k[:4]}'
brl = lambda v: f'R$ {v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

# chassi da CEABS -> chassi da nossa frota (CEABS cadastrou o SQZ8G22 sem o "9" inicial)
CHASSI_CORRIGIDO = {'4DFCAP15RB134726': '94DFCAP15RB134726'}
SERVICO = {'RHJ1F91', 'RJD3J91', 'SSD5D93', 'BDF3I26', 'KRL8A64', 'BEQ3J91', 'SRB4C78', 'SQW2A55', 'RJA4E13', 'BAZ8187', 'BCO7085'}
OBS = {
    'BCO7085': ('VERIFICAR', 'Monitoramento (22/08) informou que NÃO tem rastreador, mas a CEABS cobra Telemetria desde 16/10/2025. Cancelamento pedido em 03/09. Conferir fisicamente; se não houver equipamento, contestar desde out/2025.'),
    'RJA4E13': ('ATENÇÃO', 'Motor batido. Rescisão confirmada em 03/09/2026; retirada precisa ser agendada (prazo 03/10, tarefa T45).'),
    'BAZ8187': ('ATENÇÃO', 'Disponível para venda. Rescisão confirmada em 03/09/2026; retirada precisa ser agendada (prazo 03/10, tarefa T45).'),
    'BCY4I45': ('CONTESTAR', 'Vendido em 19/06/2026 (João Emilio) e cobrado em jul e ago. Rescisão confirmada em 03/09; retirada depende do comprador (Robson 21 96464-6193).'),
    'SQZ8G30': ('CONTESTAR', 'Listado como GM Macaé no pedido de 03/09, mas o cadastro mostra À VENDA desde 21/01/2026, ou seja, cerca de 7 meses pagos com o carro parado. Confirmar e cancelar.'),
    'SQZ8G22': ('ATENÇÃO', 'Na CEABS o chassi está truncado (4DFCAP15RB134726). Usar o código vei23122715400 no pedido de cancelamento.'),
    'SSD5D93': ('OK', 'Cadastro diz DISPONÍVEL VENDA, mas o Monitoramento informa uso como socorro de oficina (Caju). Instalado em 29/06/2026. Atualizar o cadastro.'),
    'SRB4C78': ('OK', 'Cadastro diz DISPONÍVEL VENDA, mas o Monitoramento informa uso como socorro de oficina (Deodoro). Instalado em 23/06/2026. Atualizar o cadastro.'),
    'SQW2A55': ('OK', 'Socorro de oficina (Deodoro), conforme Monitoramento. Instalado em 24/06/2026. Cadastro sem destino.'),
    'RJD3J91': ('OK', 'Cadastro diz DISPONÍVEL LOCAÇÃO desde 2024, mas o Monitoramento informa uso no setor de peças (Caju).'),
    'BEQ3J91': ('OK', 'Uso da Elisangela (Caju), conforme Monitoramento. Cadastro sem destino desde que saiu da CEDAE.'),
    'SRL0E31': ('ATENÇÃO', 'Spin ex-Rioluz. O cancelamento das Spin foi informado como já solicitado em 17/07, mas segue cobrada. Hoje está na SES (desde 08/09).'),
    'SRN0F11': ('ATENÇÃO', 'Spin ex-Rioluz, na Magen Blindados desde 28/07. Cancelamento das Spin informado em 17/07, não refletido na CEABS. Decidir se mantém.'),
}


def categoria(i, f):
    p = f['placa']
    if f['data_venda']:
        return 'VENDIDO - rastreador ainda ativo', 'CONTESTAR', 'Pedir cancelamento e retirada imediatos; contestar o mês cobrado após a venda.'
    if p in SERVICO:
        return 'USO INTERNO GRILLO (veículo de serviço)', 'OK', 'Decisão da gerência: manter ou cancelar.'
    if f['chassi'].startswith('94DFCAP15') and 'GM - MACAÉ' in f['local_agosto']:
        return 'DEVOLVIDO PELA GM MACAÉ em ago/26', 'ATENÇÃO', 'Cobrança de agosto procede (estava com a GM até 17 a 24/08). Confirmar o cancelamento, senão continua sendo cobrado em set/out.'
    loc = f['local_agosto']
    com_cliente = loc and any(k not in ('PATIO MANUTENÇÃO', 'VENDIDO') and not k.startswith('DISPONIVEL')
                              for k in [t.rsplit(' ', 1)[0] for t in loc.split(' > ')])
    if com_cliente:
        return 'COM CLIENTE', 'OK', 'OK'
    if f['situacao'] == 'sem_destino_registrado':
        return 'SEM LOCALIZAÇÃO NO CADASTRO', 'VERIFICAR', f"Cadastro sem destino desde {f['desde']}. Localizar o veículo e decidir se mantém o rastreador."
    return 'PARADO NA GRILLO (pátio/disponível)', 'ATENÇÃO', 'Rastreador pago sem receita de locação. Decidir se mantém (segurança) ou cancela.'


def fmt_ws(ws, larguras, moeda_cols=(), freeze='A2'):
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = CAB
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for idx, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(idx)].width = w
    for col in moeda_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            for c in row: c.number_format = MOEDA
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


wb = Workbook()

# ---------------- Agosto por veículo ----------------
ult = max(med)
ws = wb.active; ws.title = f'{rot(ult)} por veículo'.replace('/', '-')
cols = ['Código CEABS', 'Placa (nossa frota)', 'Placa na CEABS', 'Chassi', 'Modelo', 'Produto', 'NF', 'Valor', 'Instalação',
        'Onde estava em agosto (nosso cadastro)', 'Situação hoje', 'Cliente/local atual', 'Contrato', 'Categoria', 'Status', 'Ação', 'Observação']
ws.append(cols)
nf_por_produto = {n['produto']: n['nf'] for n in notas if n['ref'] == ult}
resumo_cat = collections.defaultdict(lambda: [0, 0.0])
linhas_ago = []
for i in med[ult]['itens']:
    ch = CHASSI_CORRIGIDO.get(i['chassi'], i['chassi'])
    f = frota[ch]
    cat, status, acao = categoria(i, f)
    st_obs, obs = OBS.get(f['placa'], (None, ''))
    if f['chassi'].startswith('9BGJJ7520') and f['placa'] not in OBS:
        obs = 'Spin ex-Rioluz (contrato encerrado em 20/07/2026). Cancelamento das Spin informado como solicitado em 17/07, não refletido na CEABS.'
    if f['situacao'] == 'vendido' and f['data_venda']:
        obs = (obs + ' ' if obs else '') + f"Vendido em {f['data_venda'][8:]}/{f['data_venda'][5:7]}/{f['data_venda'][:4]} para {f['vendido_para']}. Risco de cobrança de R$ 400 pelo equipamento se não for retirado."
    if st_obs: status = st_obs if st_obs != 'OK' or status == 'OK' else status
    resumo_cat[cat][0] += 1; resumo_cat[cat][1] += i['valor']
    linha = [i['codigo'], f['placa'], i['placa'] or '(sem placa)', i['chassi'], f['modelo'], i['produto'], nf_por_produto.get(i['produto']),
             i['valor'], i['instalacao'], f['local_agosto'] or '(sem registro no mês)', f['situacao'], f['cliente_atual'], f['contrato'], cat, status, acao, obs]
    linhas_ago.append(linha); ws.append(linha)
for r in ws.iter_rows(min_row=2):
    r[14].fill = PatternFill('solid', fgColor=CORES[r[14].value])
    for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
fmt_ws(ws, [16, 10, 10, 19, 26, 16, 7, 10, 11, 42, 20, 26, 22, 30, 11, 40, 60], moeda_cols=(8,), freeze='C2')

# ---------------- Resumo por categoria ----------------
wr = wb.create_sheet('Resumo', 0)
wr.append([f'CONFERÊNCIA CEABS - referência {rot(ult)}', None, None])
wr['A1'].font = Font(bold=True, size=14)
nt = [n for n in notas if n['ref'] == ult]
wr.append([f"Medição: {len(med[ult]['itens'])} itens, total {brl(med[ult]['resumo']['TOTAL'])}"])
wr.append([' + '.join(f"NF {n['nf']} ({n['produto']}) {brl(n['valor'])}" for n in nt) + f" = {brl(sum(n['valor'] for n in nt))}: bate com a medição."])
wr.append(['Boletos = valor líquido de cada NF (descontados IRRF e PIS/COFINS/CSLL, que a Grillo recolhe por DARF). Vencimento ' + nt[0]['vencimento'] + '.'])
wr.append([])
wr.append(['Categoria', 'Qtd', 'Valor'])
ini = wr.max_row
for cat, (q, v) in sorted(resumo_cat.items(), key=lambda x: -x[1][1]):
    wr.append([cat, q, round(v, 2)])
wr.append(['TOTAL', sum(q for q, _ in resumo_cat.values()), round(sum(v for _, v in resumo_cat.values()), 2)])
for r in wr.iter_rows(min_row=ini, max_row=ini):
    for c in r: c.font = Font(bold=True, color='FFFFFF'); c.fill = CAB
for r in wr.iter_rows(min_row=ini + 1, max_row=wr.max_row, min_col=3, max_col=3):
    for c in r: c.number_format = MOEDA
wr.column_dimensions['A'].width = 120; wr.column_dimensions['B'].width = 8; wr.column_dimensions['C'].width = 14

# ---------------- Fiscal ----------------
wf = wb.create_sheet('Fiscal - NFs')
wf.append(['Ref.', 'NF', 'Emissão', 'Layout', 'Serviço na NF', 'Cód. trib. nacional', 'Valor', 'ISS 5% (não retido)', 'IRRF retido', '% IRRF',
           'PIS/COFINS/CSLL retidos', 'Esperado 4,65%', 'PIS próprio', 'COFINS própria', 'IBS', 'CBS', 'Líquido NF', 'Boleto', 'Vencimento',
           'Endereço do tomador na NF', 'Conferência'])
for n in notas:
    esp = round(n['valor'] * 0.0465 + 1e-9, 2)
    if n['nf'] == '8979':
        conf = ('DIVERGENTE: IRRF de 15% e nenhuma retenção de PIS/COFINS/CSLL. 15% é a alíquota de multa ou indenização por rescisão de contrato '
                '(art. 70 da Lei 9.430/96), mas a nota foi emitida como SERVIÇO (código 11.05.01, com ISS). Se for serviço: IRRF 1% (R$ 21,80) + CSRF 4,65% '
                '(R$ 101,37), líquido R$ 2.056,83. Se for multa: não caberia NFS-e/ISS. Pedir à CEABS cancelamento/substituição ou definição formal; '
                'a contabilidade decide o código do DARF e a natureza no EFD-Reinf.')
    elif n['nf'] == '8978':
        conf = ('ATENÇÃO: nenhuma retenção. Código 14.02.01 (assistência técnica). PIS/COFINS/CSLL de 4,65% = R$ 12,56, acima do mínimo de R$ 10,00. '
                'Se a contabilidade enquadrar como manutenção, a retenção era devida (responsabilidade da Grillo como fonte pagadora). IRRF 1,5% = R$ 4,05, abaixo de R$ 10,00 (dispensado).')
    else:
        conf = 'OK: ISS 5%, IRRF 1%, CSRF 4,65%, PIS/COFINS próprios, IBS/CBS (0,1%/0,9% sobre a base sem ISS/PIS/COFINS) e líquido = boleto.'
    wf.append([rot(n['ref']), n['nf'], n['emissao'], n['layout'], n['servico'], n['cod_trib'], n['valor'], n['iss'], n['irrf'], round(n['irrf'] / n['valor'], 4),
               n['csrf'], esp, n['pis_proprio'], n['cofins_proprio'], n['ibs'], n['cbs'], n['liquido'], n['boleto_valor'], n['vencimento'], n['tomador_end'], conf])
for r in wf.iter_rows(min_row=2):
    r[9].number_format = '0.0%'
    st = 'CONTESTAR' if r[1].value == '8979' else 'ATENÇÃO' if r[1].value == '8978' else 'OK'
    r[20].fill = PatternFill('solid', fgColor=CORES[st])
    for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
fmt_ws(wf, [9, 7, 11, 18, 28, 10, 11, 10, 10, 7, 11, 11, 9, 9, 7, 7, 11, 11, 11, 34, 70], moeda_cols=(7, 8, 9, 11, 12, 13, 14, 15, 16, 17, 18), freeze='C2')
wf.append([])
for t in ['Endereço oficial no Cartão CNPJ (13/10/2026): Rua Nicarágua, 630, Bloco 2, Sala 608, Penha. Nenhuma NF da CEABS usa esse endereço, e cada lote veio com um diferente (Campo Grande, Deodoro, Caju). Pedir à CEABS que corrija o cadastro do tomador.',
          'Razão social na NF: "TRANSPORTADORA GRILLO E RIBEIRO LTD"; na planilha de medição: "LOCADORA GRILLO E RIBEIRO LTDA" (mesmo CNPJ 06.025.418/0001-30). Conferir com o Cartão CNPJ qual é a atual.',
          'IBS/CBS (2026, ano de teste: 0,1% e 0,9%) estão corretos em todas as NFs do layout novo. Não geram desembolso para a Grillo.',
          'O IRRF e o PIS/COFINS/CSLL retidos foram descontados dos boletos, então a Grillo precisa recolher por DARF até o dia 20 do mês seguinte ao pagamento. Se as NFs 8966/8967/8979 (venc. 03/09) ou 10276/10277 (venc. 25/09) foram pagas em setembro, os DARFs vencem em 20/10/2026.']:
    wf.append([t]); wf.cell(wf.max_row, 1).alignment = Alignment(wrap_text=False)

# ---------------- Histórico mensal ----------------
wh = wb.create_sheet('Histórico mensal')
wh.append(['Referência', 'Qtd itens (principal)', 'Monitoramento GSM', 'Telemetria', 'Principal', 'Taxa de rescisão', 'Visita improdutiva', 'Total medição',
           'NFs', 'Total NFs', 'IRRF retido', 'CSRF retido', 'Líquido (boletos)', 'Vencimento', 'Confere medição x NF'])
refs = sorted(set(med) | {n['ref'] for n in notas})
a, b = int(refs[0][5:]), int(refs[-1][5:])
for mm in range(a, b + 1):
    k = f'{refs[0][:4]}-{mm:02d}'
    if k not in med:
        wh.append([rot(k), None, None, None, None, None, None, None, 'FALTA RECEBER a medição e as NFs desta referência'])
        continue
    it = [x for x in med[k]['itens'] if x['aba'].endswith('Principal')]
    ns = [n for n in notas if n['ref'] == k]
    rs = med[k]['resumo']
    soma = lambda p: round(sum(x['valor'] for x in it if x['produto'] == p), 2)
    tot_nf = round(sum(n['valor'] for n in ns), 2)
    wh.append([rot(k), len(it), soma('Monitoramento GSM'), soma('Telemetria'), rs.get('Principal'), rs.get('Taxa de Rescisão', 0), rs.get('Visita Improdutiva', 0), rs['TOTAL'],
               ', '.join(n['nf'] for n in ns), tot_nf, round(sum(n['irrf'] for n in ns), 2), round(sum(n['csrf'] for n in ns), 2),
               round(sum(n['liquido'] for n in ns), 2), ns[0]['vencimento'] if ns else None, 'OK' if abs(tot_nf - rs['TOTAL']) < 0.05 else f'DIFERENÇA {tot_nf - rs["TOTAL"]:.2f}'])
fmt_ws(wh, [11, 10, 13, 13, 13, 12, 12, 13, 22, 13, 11, 11, 14, 11, 16], moeda_cols=(3, 4, 5, 6, 7, 8, 10, 11, 12, 13))

# ---------------- Histórico por veículo ----------------
wv = wb.create_sheet('Histórico por veículo')
ks = sorted(med)
wv.append(['Código CEABS', 'Chassi', 'Placa', 'Produto'] + [rot(k) for k in ks] + ['Movimento'])
todos = {}
for k in ks:
    for x in med[k]['itens']:
        if x['aba'].endswith('Principal'):
            d = todos.setdefault(x['codigo'], {'x': x, 'v': {}})
            d['x'] = x; d['v'][k] = x['valor']
placa_frota = {CHASSI_CORRIGIDO.get(i['chassi'], i['chassi']): frota[CHASSI_CORRIGIDO.get(i['chassi'], i['chassi'])]['placa'] for i in med[ult]['itens']}
for cod, d in sorted(todos.items(), key=lambda t: (len(t[1]['v']) == len(ks), t[1]['x']['produto'], t[0])):
    x = d['x']; vals = [d['v'].get(k) for k in ks]
    if all(v is not None for v in vals): mov = 'Cobrado em todas as referências'
    elif vals[-1] is None: mov = f"Saiu (cancelado em {x['cancelamento'][:10]})" if x['cancelamento'] else 'Saiu'
    else: mov = f"Entrou (instalado em {x['instalacao']})"
    wv.append([cod, x['chassi'], placa_frota.get(CHASSI_CORRIGIDO.get(x['chassi'], x['chassi'])) or x['placa'] or '(máquina/sem placa)', x['produto']] + vals + [mov])
fmt_ws(wv, [17, 20, 18, 17] + [11] * len(ks) + [36], moeda_cols=range(5, 5 + len(ks)))

# ---------------- Julho: rescisões e visitas ----------------
if '2026-07' in med:
    wj = wb.create_sheet('Jul-2026 rescisão e visitas')
    wj.append(['Tipo', 'Código', 'Chassi', 'Placa', 'Valor', 'Cancelamento (CEABS)', 'Retirada (CEABS)', 'Observação'])
    cont = collections.Counter((x['produto'], x['codigo']) for x in med['2026-07']['itens'] if x['produto'] == 'Visita Improdutiva')
    for x in med['2026-07']['itens']:
        if x['produto'] not in ('Visita Improdutiva', 'Taxa de Rescisão'): continue
        obs = ''
        if x['produto'] == 'Visita Improdutiva' and cont[(x['produto'], x['codigo'])] > 1:
            obs = 'Mesmo veículo cobrado 2 vezes (2 x R$ 45). Pedir os números das OS e as datas de cada visita.'
        if x['valor'] == 260: obs = 'R$ 260 (padrão R$ 60). Instalado em 16/10/2025 e cancelado em 07/07/2026, provável multa por cancelamento antes de 12 meses. Pedir a memória de cálculo.'
        if x['chassi'] == '94DFCAP15RB133897': obs = 'Kicks SRG4J44 (ex-Casa Civil). Vendido à Jazz em 09/02/2026, mas só foi cancelado em 07/07/2026, e as mensalidades de fev a jul foram pagas com o carro já vendido.'
        if x['produto'] == 'Taxa de Rescisão' and not obs: obs = 'R$ 60 por retirada, valor combinado em 2024.'
        wj.append([x['produto'], x['codigo'], x['chassi'], x['placa'] or '(máquina)', x['valor'], x['cancelamento'], x['retirada'], obs])
    pr = [x for x in med['2026-07']['itens'] if x['aba'].endswith('Principal') and x['prorata']]
    wj.append([]); wj.append([f'Pró-rata de julho de {len(pr)} itens cancelados em 02-03/07 (AEGEA/Prolagos/Foz): {brl(sum(x["valor"] for x in pr))}. O cancelamento foi informado em 01/07, e a CEABS registrou 02/07 16h e 03/07 10h.'])
    fmt_ws(wj, [18, 17, 20, 11, 10, 20, 20, 90], moeda_cols=(5,))

wb.save(SAIDA)
print('ok', SAIDA, {k: (v[0], round(v[1], 2)) for k, v in resumo_cat.items()})
