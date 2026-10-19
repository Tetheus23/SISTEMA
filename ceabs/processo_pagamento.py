"""Processo de pagamento da CEABS no padrão do processo da Ticket Log.

Uso: python3 -I processo_pagamento.py <historico.xlsx> <forpag_modelo.xlsb> <mes AAAA-MM> <saida_dir> <nf_ou_boleto.pdf>...
  historico.xlsx   planilha de historico_cobrancas.py (de preferência a revisada: abas Painel e Histórico)
  forpag_modelo    modelo do FORPAG (aba FORMULÁRIO)
  pdfs             NFS-e (DANFSe) e boletos do mês, em qualquer ordem: o tipo é lido do texto e o boleto casa com a nota
                   pelo valor líquido
Gera em <saida_dir>:
  FORPAG CEABS - <MES>.xlsx / .pdf      formulário de liberação preenchido
  PROCESSO PAGAMENTO CEABS - <MES>.pdf  1 FORPAG, 2 notas fiscais e boletos, 3 relatório resumido, 4 consolidado
  EMAIL DIRETORIA - CEABS <MES>.html    pedido de autorização de pagamento, com o resumo por cliente
Precisa do LibreOffice (uno), do Chromium do Playwright e de pypdf/reportlab.
"""
import datetime as dt
import glob
import html
import io
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from collections import Counter, defaultdict

import openpyxl
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen import canvas

HIST, MODELO, MES, SAIDA = sys.argv[1:5]
ANEXOS = sys.argv[5:]
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
ANO, NMES = int(MES[:4]), int(MES[5:7])
ROTULO = f'{MESES[NMES - 1].capitalize()}/{MES[2:4]}'           # Ago/26
SUFIXO = f'{MESES[NMES - 1].upper()}-{MES[2:4]}'                 # AGO-26
CHROME = os.environ.get('CHROME') or (glob.glob('/opt/pw-browsers/chromium-*/chrome-linux/chrome') or ['chromium'])[0]
EMPRESA = 'LOCADORA GRILLO E RIBEIRO LTDA'
TIMBRE = ('CNPJ 06.025.418/0001-30<br>Rua Nicarágua, nº 630, Bloco 2, Sala 608, Penha, Rio de Janeiro/RJ, CEP 21.020-050<br>'
          '(21) 3462-7196 | adm06@grillolocacao.com.br')
STATUS = ['MANTER', 'PEDIR CANCELAMENTO', 'JÁ SOLICITADO CANCELAMENTO', 'CANCELADO E RETIRADO',
          'CANCELADO E RETIRADA A CONFIRMAR']
COR_STATUS = dict(zip(STATUS, ['#C6EFCE', '#F4B6B6', '#FFE699', '#D9D9D9', '#F8CBAD']))


def brl(v, sinal=False):
    if v is None or v == '':
        return ''
    s = f'{abs(v):,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')
    return ('-' if v < 0 else ('+' if sinal and v > 0 else '')) + f'R$ {s}'


def num_br(s):
    return float(s.replace('.', '').replace(',', '.'))


def data(v):
    return v.strftime('%d/%m/%Y') if hasattr(v, 'strftime') else (v or '')


def mes_txt(v):
    return f'{MESES[v.month - 1]}/{v.strftime("%y")}' if hasattr(v, 'month') else ''


def e(v):
    return html.escape(str(v if v is not None else ''))


def texto_pdf(p):
    return subprocess.run(['pdftotext', '-layout', p, '-'], capture_output=True, text=True, check=True).stdout


# ---------- planilha: histórico do mês e painel ----------
wb = openpyxl.load_workbook(HIST, data_only=True)
wsh = wb['Histórico']
cab = [c.value for c in wsh[2]]
hist = [dict(zip(cab, r)) for r in wsh.iter_rows(min_row=3, values_only=True) if r[3]]
ref = dt.datetime(ANO, NMES, 1)
itens = [r for r in hist if r['MÊS/ANO APURADO'] == ref]
if not itens:
    raise SystemExit(f'nenhum item com MÊS/ANO APURADO {MES} em {HIST}')
anterior = max((r['MÊS/ANO APURADO'] for r in hist if r['MÊS/ANO APURADO'] and r['MÊS/ANO APURADO'] < ref), default=None)
itens_ant = [r for r in hist if r['MÊS/ANO APURADO'] == anterior]


def secoes(ws):
    """Tabelas do Painel por número de seção: cabeçalho 2 linhas abaixo do título, linhas até a primeira vazia."""
    lin = [[ws.cell(r, c).value for c in range(1, 11)] for r in range(1, ws.max_row + 1)]
    out = {}
    for i, x in enumerate(lin):
        m = re.match(r'(\d)\. ', str(x[0] or ''))
        if not m:
            continue
        if m[1] == '5':
            out[5] = []
            for y in lin[i + 1:]:
                if not y[0]:
                    break
                out[5].append(str(y[0]).lstrip('- '))
            continue
        titulos = [t for t in lin[i + 2] if t]
        dados = []
        for y in lin[i + 3:]:
            if y[0] in (None, 'Total'):
                break
            dados.append(y[:len(titulos)])
        out[int(m[1])] = (titulos, dados)
    return out


painel = secoes(wb['Painel'])

# ---------- notas e boletos ----------
notas, boletos = {}, []
for p in ANEXOS:
    t = texto_pdf(p)
    if 'DANFSe' in t or 'NÚMERO DA NFS-E' in t:
        g = lambda rot, linhas=1: re.search(rot + r'[^\n]*\n' + r'[^\n]*\n' * (linhas - 1) + r'\s*([\d./:, ]+?)\s{2,}', t)
        numero = re.search(r'NÚMERO DA NFS-E[^\n]*\n\s*(\d+)\s+(\S+)\s+(\S+)', t)
        ret = re.search(r'IRRF\s+Contribuição Previdenciária - Retida\s+Contribuições Sociais - Retidas\s*\n\s*([\d.,]+)\s+\S+\s+([\d.,]+)', t)
        tot = re.search(r'Total das Retenções \(ISSQN/Federais\)\s+VALOR LÍQUIDO DA NFS-E[^\n]*\n\s*([\d.,]+)\s+([\d.,]+)', t)
        bruto = re.search(r'VALOR DA OPERAÇÃO / SERVIÇO[^\n]*\n\s*([\d.,]+)', t)
        servico = re.search(r'IBS e CBS\.,\s*([^\n]+)', t)
        notas[numero[1]] = dict(arquivo=p, numero=numero[1], emissao=numero[3], bruto=num_br(bruto[1]), irrf=num_br(ret[1]),
                                csrf=num_br(ret[2]), retencoes=num_br(tot[1]), liquido=num_br(tot[2]),
                                servico=(servico[1].strip() if servico else '').title())
    elif 'Recibo do Sacado' in t or 'Valor do documento' in t:
        valor = re.search(r'Uso do Banco\s*\n\s*([\d.]+,\d{2})', t) or re.search(r'Valor do documento.*?([\d.]+,\d{2})', t, re.S)
        venc = re.search(r'Vencimento.*?(\d{2}/\d{2}/\d{4})', t, re.S)
        boletos.append(dict(arquivo=p, valor=num_br(valor[1]), vencimento=venc[1]))
    else:
        raise SystemExit(f'não reconheci {p} como NFS-e nem boleto')
for b in boletos:
    n = next((n for n in notas.values() if abs(n['liquido'] - b['valor']) < 0.005 and 'boleto' not in n), None)
    if not n:
        raise SystemExit(f'boleto de {brl(b["valor"])} não casa com nenhuma nota')
    n['boleto'] = b
ordem = sorted(notas.values(), key=lambda n: int(n['numero']))
total_pagar = round(sum(n['boleto']['valor'] if 'boleto' in n else n['liquido'] for n in ordem), 2)
vencimento = min(n['boleto']['vencimento'] if 'boleto' in n else '' for n in ordem)
nfs_txt = ' e '.join(n['numero'] for n in ordem)
print('notas:', [(n['numero'], n['bruto'], n['liquido'], n.get('boleto', {}).get('valor')) for n in ordem], '| a pagar', total_pagar)

# ---------- 1. FORPAG (LibreOffice) ----------
os.makedirs(SAIDA, exist_ok=True)
FORPAG_XLSX = os.path.join(SAIDA, f'FORPAG CEABS - {SUFIXO}.xlsx')
FORPAG_PDF = os.path.join(SAIDA, f'FORPAG CEABS - {SUFIXO}.pdf')
campos = {'H21': total_pagar, 'D22': ' / '.join(n['numero'] for n in ordem), 'D23': vencimento,
          'H23': '=H21-H22',  # o modelo traz =H21-H2
          'G61': f'CEABS - RASTREADORES\nMEDIÇÃO {ROTULO.upper()} - NFS-e {nfs_txt}\nVENC. {vencimento}',
          # cargos das assinaturas: o modelo usa XLOOKUP, que o LibreOffice não calcula (#MACRO?)
          'C58': 'Ger. Oficina', 'E58': 'Coord. Frota', 'E65': 'Coord. Frota'}


def preenche_forpag():
    import uno
    from com.sun.star.beans import PropertyValue

    def prop(nome, valor):
        p = PropertyValue()
        p.Name, p.Value = nome, valor
        return p
    porta = 2900 + os.getpid() % 600
    perfil = tempfile.mkdtemp(prefix='lo_forpag_')
    soffice = subprocess.Popen(['soffice', '--headless', '--invisible', '--norestore', '--nologo',
                                f'-env:UserInstallation=file://{perfil}', f'--accept=socket,host=127.0.0.1,port={porta};urp;'],
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                               env=dict(os.environ, LANG='pt_BR.UTF-8', LC_ALL='pt_BR.UTF-8'))  # R$ 4.099,69
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
        for _ in range(60):
            try:
                ctx = resolver.resolve(f'uno:socket,host=127.0.0.1,port={porta};urp;StarOffice.ComponentContext')
                break
            except Exception:
                time.sleep(1)
        else:
            raise SystemExit('LibreOffice não respondeu')
        desktop = ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', ctx)
        copia = os.path.join(perfil, 'modelo' + os.path.splitext(MODELO)[1])
        shutil.copy(MODELO, copia)
        doc = desktop.loadComponentFromURL(uno.systemPathToFileUrl(copia), '_blank', 0, (prop('Hidden', True),))
        if doc is None:
            raise SystemExit(f'LibreOffice não abriu {MODELO}')
        aba = doc.Sheets.getByName('FORMULÁRIO')
        for end, v in campos.items():
            cel = aba.getCellRangeByName(end)
            if isinstance(v, (int, float)):
                cel.Value = v
            elif v.startswith('='):
                cel.Formula = v
            else:
                cel.String = v
        # uma página: área do formulário (sem o botão LIMPAR CAMPOS), ajustada à folha
        aba.setPrintAreas((aba.getCellRangeByName('C3:L70').RangeAddress,))
        estilo = doc.StyleFamilies.getByName('PageStyles').getByName(aba.PageStyle)
        estilo.ScaleToPagesX, estilo.ScaleToPagesY, estilo.CenterHorizontally = 1, 1, True
        doc.calculateAll()
        doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(FORPAG_XLSX)), (prop('FilterName', 'Calc MS Excel 2007 XML'),))
        for i in range(doc.Sheets.Count):  # PDF só do formulário
            s = doc.Sheets.getByIndex(i)
            if s.Name != 'FORMULÁRIO':
                s.IsVisible = False
        doc.storeToURL(uno.systemPathToFileUrl(os.path.abspath(FORPAG_PDF)), (prop('FilterName', 'calc_pdf_Export'),))
        doc.close(True)
    finally:
        soffice.terminate()
        try:
            soffice.wait(20)
        except subprocess.TimeoutExpired:
            soffice.kill()


preenche_forpag()
if len(PdfReader(FORPAG_PDF).pages) != 1:
    print('AVISO: o FORPAG saiu com', len(PdfReader(FORPAG_PDF).pages), 'páginas')

# ---------- 3 e 4. relatório (HTML -> PDF no Chromium) ----------
CSS = """
@page { size: A4; margin: 11mm 11mm 15mm 11mm; }
@page paisagem { size: A4 landscape; }
* { box-sizing: border-box; }
body { font-family: Carlito, 'Liberation Sans', sans-serif; color: #1b1b1b; font-size: 9.5pt; margin: 0; }
.topo { display: flex; justify-content: space-between; align-items: flex-end; border-bottom: 2px solid #1F3A5F; padding-bottom: 5px; }
.topo b { font-size: 12pt; color: #1F3A5F; }
.topo div { font-size: 7.5pt; color: #444; line-height: 1.35; }
.topo .dir { text-align: right; }
.secao { color: #2a78d6; font-weight: bold; font-size: 8.5pt; letter-spacing: .4px; margin: 10px 0 2px; }
h1 { font-size: 17pt; color: #1F3A5F; margin: 0; }
.sub { color: #555; margin: 1px 0 8px; }
h2 { font-size: 11.5pt; color: #1F3A5F; border-left: 4px solid #2a78d6; padding-left: 6px; margin: 12px 0 5px; }
.cards { display: grid; grid-template-columns: repeat(3, 1fr); gap: 6px; margin-bottom: 4px; }
.duas { display: grid; grid-template-columns: 1.15fr 1fr; gap: 12px; }
tr.grupo td { border-top: 1px solid #9BB7D4; }
.card { border: 1px solid #d6dde6; border-radius: 4px; padding: 6px 8px; }
.card.forte { background: #1F3A5F; color: #fff; border-color: #1F3A5F; }
.card .r { font-size: 7pt; text-transform: uppercase; color: #6b7280; letter-spacing: .3px; }
.card.forte .r, .card.forte .n { color: #dbe6f3; }
.card .v { font-size: 15pt; font-weight: bold; margin: 1px 0; }
.card .n { font-size: 7.5pt; color: #555; }
table { border-collapse: collapse; width: 100%; font-size: 8.5pt; }
th { background: #1F3A5F; color: #fff; text-align: left; padding: 3px 5px; font-weight: bold; }
td { padding: 2.5px 5px; border-bottom: 1px solid #e3e6ea; vertical-align: top; }
td.n, th.n { text-align: right; white-space: nowrap; }
tr.tot td { background: #DDEBF7; font-weight: bold; border-top: 1px solid #9BB7D4; }
tr.sub td { background: #F2F2F2; font-weight: bold; }
thead { display: table-header-group; }
tr { page-break-inside: avoid; }
.nota { color: #666; font-size: 7.5pt; margin-top: 3px; }
.aviso { border: 1px solid #f0c36d; background: #fff8e6; border-radius: 4px; padding: 6px 9px; margin-top: 6px; }
.aviso li { margin: 2px 0; }
.ok { color: #0a7a33; font-weight: bold; }
.st { padding: 0 4px; border-radius: 3px; white-space: nowrap; font-size: 7.5pt; }
.quebra { page-break-before: always; }
.consol { page: paisagem; }
.consol table { font-size: 7.5pt; }
.consol td { white-space: nowrap; padding: 2px 4px; }
.consol th { white-space: nowrap; }
"""


def topo():
    return (f'<div class="topo"><div><b>{EMPRESA}</b><br>{TIMBRE}</div>'
            f'<div class="dir">Processo de pagamento<br>CEABS · contrato 40010217</div></div>')


def tabela(titulos, linhas, num=(), classes=None, status_col=None):
    h = '<table><thead><tr>' + ''.join(f'<th class="{"n" if i in num else ""}">{e(t)}</th>' for i, t in enumerate(titulos)) + '</tr></thead><tbody>'
    for k, lin in enumerate(linhas):
        cl = (classes or {}).get(k, '')
        h += f'<tr class="{cl}">'
        for i, v in enumerate(lin):
            if i == status_col and v in COR_STATUS:
                h += f'<td><span class="st" style="background:{COR_STATUS[v]}">{e(v)}</span></td>'
            else:
                h += f'<td class="{"n" if i in num else ""}">{v if str(v).startswith("<") else e(v)}</td>'
        h += '</tr>'
    return h + '</tbody></table>'


def card(rotulo, valor, nota, forte=False):
    return f'<div class="card{" forte" if forte else ""}"><div class="r">{e(rotulo)}</div><div class="v">{e(valor)}</div><div class="n">{e(nota)}</div></div>'


soma = lambda rs, k: round(sum(r[k] or 0 for r in rs), 2)
cheio, liquido = soma(itens, 'Valor a Faturar'), soma(itens, 'VALOR SEM IMPOSTO')
MENSAL = lambda r: r['Produto'] not in ('Taxa de Rescisão', 'Visita Improdutiva')
n_equip = len({r['Código'] for r in itens if MENSAL(r)})
n_equip_ant = len({r['Código'] for r in itens_ant if MENSAL(r)})
liq_ant = soma([r for r in itens_ant if MENSAL(r)], 'VALOR SEM IMPOSTO')  # sem taxas e visitas (contestadas em jul/26)
unit = {}  # mensalidade cheia de cada produto, com a retenção do boleto
for p in sorted({r['Produto'] for r in itens if MENSAL(r)}):
    rs = [r for r in itens if r['Produto'] == p]
    cheia = Counter(r['Valor a Faturar'] for r in rs).most_common(1)[0][0]
    unit[p] = round(cheia * soma(rs, 'VALOR SEM IMPOSTO') / soma(rs, 'Valor a Faturar'), 2)

# conferência medição x nota x boleto (só no console; o relatório mostra o resultado)
por_nf = defaultdict(list)
for r in itens:
    por_nf[r['NOTAS FISCAIS']].append(r)
for n in ordem:
    rs = next((v for k, v in por_nf.items() if re.search(rf'\b{n["numero"]}\b', k or '')), [])
    bol = n.get('boleto', {}).get('valor')
    if abs(soma(rs, 'Valor a Faturar') - n['bruto']) >= 0.01 or bol is None or abs(bol - n['liquido']) >= 0.01:
        print(f'AVISO: NFS-e {n["numero"]}: medição {soma(rs, "Valor a Faturar")}, nota {n["bruto"]}/{n["liquido"]}, boleto {bol}')

# 3. relatório resumido
_, casos = painel[4]
modelo = {x[1]: x[2] for x in casos}
pedido_em = {x[1]: (re.search(r'Pedido em (\S+?):', x[6] or '') or [None, ''])[1] for x in casos}
veic = {}
for r in itens:
    veic[r['Placa']] = r
pedir = sorted(p for p, r in veic.items() if r['STATUS'] == 'PEDIR CANCELAMENTO')
solicitado = sorted((p for p, r in veic.items() if r['STATUS'] == 'JÁ SOLICITADO CANCELAMENTO'), key=lambda p: (pedido_em.get(p, ''), p))
cli = defaultdict(Counter)  # cliente -> status -> veículos (pode estar cobrando, mas já cancelado ou a cancelar)
for r in veic.values():
    cli[r['CLIENTE NO PERÍODO DA MEDIÇÃO'] or '(sem registro)'][r['STATUS']] += 1
lin_cli, grupo_cli = [], {}
for c, sts in sorted(cli.items(), key=lambda kv: (-sum(kv[1].values()), kv[0])):
    grupo_cli[len(lin_cli)] = 'grupo'
    for k, (st, n) in enumerate(sorted(sts.items(), key=lambda kv: STATUS.index(kv[0]))):
        lin_cli.append([c if k == 0 else '', st, n])
grupo_cli[len(lin_cli)] = 'tot'
lin_cli.append(['Total', '', len(veic)])
var = lambda a, b: f'{(a / b - 1) * 100:+.1f}%'.replace('.', ',') if b else ''
# histórico de pagamentos: 3 meses anteriores (notas pagas; as contestadas ficam de fora) e o mês deste processo
_, meses_painel = painel[2]   # mês apurado, vencimento, notas, equipamentos, ...
_, notas_painel = painel[3]   # mês apurado, nota, tipo, vencimento, bruto, líquido, pagamento, situação
veic_mes = {x[0]: x[3] for x in meses_painel}
lin_pag, contestadas = [], []
for m in sorted({x[0] for x in notas_painel if x[0] < ref})[-3:]:
    ns = [x for x in notas_painel if x[0] == m and not str(x[7]).startswith('Contestado')]
    contestadas += [x for x in notas_painel if x[0] == m and str(x[7]).startswith('Contestado')]
    pagos = [x[6] for x in ns]
    lin_pag.append([mes_txt(m), mes_txt(min(x[3] for x in ns)), veic_mes.get(m, ''), brl(sum(x[4] for x in ns)),
                    brl(sum(x[5] for x in ns)), data(max(pagos)) if all(pagos) else 'em aberto'])
lin_pag.append([mes_txt(ref), mes_txt(dt.datetime.strptime(vencimento, '%d/%m/%Y')), n_equip, brl(cheio), brl(total_pagar), 'a pagar'])
nota_pag = ''
if contestadas:
    nota_pag = (f'Fora do histórico: {" e ".join(x[1] for x in contestadas)} ({", ".join(sorted({str(x[2]).lower() for x in contestadas}))}, '
                f'{brl(sum(x[4] for x in contestadas))}), contestadas com a CEABS e não pagas.')
h3 = f"""{topo()}
<div class="secao">3 · RELATÓRIO RESUMIDO</div>
<h1>CEABS · Rastreadores · {ROTULO}</h1>
<div class="sub">Medição de {MESES[NMES - 1]}/{ANO} · NFS-e {nfs_txt} · vencimento {vencimento}</div>
<div class="cards">
{card('Equipamentos instalados', n_equip, f'mês anterior ({mes_txt(anterior)}): {n_equip_ant} ({n_equip - n_equip_ant:+d})')}
{card('Valor líquido', brl(total_pagar), f'mês anterior: {brl(liq_ant)} ({var(total_pagar, liq_ant)})', True)}
{card('Valor unitário (líquido)', brl(total_pagar / n_equip), ' · '.join(f'{p}: {brl(v)}' for p, v in unit.items()))}
</div>
<div class="duas">
<div><h2>Veículos cobrados por cliente</h2>
{tabela(['Cliente no período', 'Status', 'Veículos'], lin_cli, num=(2,), classes=grupo_cli, status_col=1)}
<h2>Histórico de pagamentos</h2>
{tabela(['Mês apurado', 'Vencimento', 'Veículos', 'Valor bruto', 'Valor líquido', 'Pagamento'], lin_pag, num=(2, 3, 4),
        classes={len(lin_pag) - 1: 'tot'})}
<div class="nota">{e(nota_pag)}</div></div>
<div><h2>Solicitar cancelamento ({len(pedir)})</h2>
{tabela(['Placa', 'Modelo', 'Cliente no período'], [[p, modelo.get(p, ''), veic[p]['CLIENTE NO PERÍODO DA MEDIÇÃO']] for p in pedir])}
<h2>Cancelamento já solicitado ({len(solicitado)})</h2>
{tabela(['Placa', 'Modelo', 'Pedido em'], [[p, modelo.get(p, ''), pedido_em.get(p, '')] for p in solicitado])}
<div class="nota">Cobrados normalmente em {ROTULO.lower()}: o cancelamento foi pedido em setembro.</div></div>
</div>
"""

# 4. consolidado (paisagem)
ordem_nf = {f'NFS-e {n["numero"]}': i for i, n in enumerate(ordem)}
cons, classes = [], {}
for nf, rs in sorted(por_nf.items(), key=lambda kv: ordem_nf.get(kv[0], 99)):
    for r in sorted(rs, key=lambda r: (r['CLIENTE NO PERÍODO DA MEDIÇÃO'] or '', r['Placa'] or '')):
        cons.append([r['Placa'], r['Chassi'], r['Código'], r['Produto'], data(r['Data de Instalação']), brl(r['Valor a Faturar']),
                     brl(r['VALOR SEM IMPOSTO']), nf.replace('NFS-e ', ''), r['CLIENTE NO PERÍODO DA MEDIÇÃO'], r['CLIENTE ATUAL'], r['STATUS']])
    classes[len(cons)] = 'sub'
    cons.append([f'{nf} · {len(rs)} itens', '', '', '', '', brl(soma(rs, 'Valor a Faturar')), brl(soma(rs, 'VALOR SEM IMPOSTO')), '', '', '', ''])
classes[len(cons)] = 'tot'
cons.append([f'Total · {len(itens)} itens', '', '', '', '', brl(cheio), brl(liquido), '', '', '', ''])
h5 = f"""<div class="consol">{topo()}
<div class="secao">4 · CONSOLIDADO</div>
<h2>Itens da medição de {MESES[NMES - 1]}/{ANO} · {len(itens)} itens</h2>
{tabela(['Placa', 'Chassi', 'Código', 'Produto', 'Instalação', 'Valor a faturar', 'Valor sem imposto', 'NFS-e', 'Cliente no período',
         'Cliente atual', 'Status'], cons, num=(5, 6), classes=classes, status_col=10)}
<div class="nota">Valor sem imposto = boleto da nota rateado entre os itens, proporcional ao valor a faturar.</div></div>"""

tmp = tempfile.mkdtemp(prefix='processo_ceabs_')


def chrome_pdf(corpo, nome):
    arq = os.path.join(tmp, nome + '.html')
    open(arq, 'w', encoding='utf-8').write(f'<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><style>{CSS}</style></head><body>{corpo}</body></html>')
    pdf = os.path.join(tmp, nome + '.pdf')
    subprocess.run([CHROME, '--headless=new', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', '--run-all-compositor-stages-before-draw',
                    f'--print-to-pdf={pdf}', 'file://' + arq], check=True, capture_output=True, timeout=120)
    return pdf


rel = chrome_pdf(h3, 'relatorio')
consol = chrome_pdf(h5, 'consolidado')

# ---------- junta tudo e numera as páginas do relatório ----------
w = PdfWriter()
for p in [FORPAG_PDF] + [x for n in ordem for x in (n['arquivo'], n.get('boleto', {}).get('arquivo')) if x]:
    for pg in PdfReader(p).pages:
        w.add_page(pg)
inicio_rel = len(w.pages)
for p in (rel, consol):
    for pg in PdfReader(p).pages:
        w.add_page(pg)
total = len(w.pages)
for i in range(inicio_rel, total):
    pg = w.pages[i]
    larg, alt = float(pg.mediabox.width), float(pg.mediabox.height)
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=(larg, alt))
    c.setFont('Helvetica', 7)
    c.setFillGray(0.35)
    c.drawString(31, 20, f'{EMPRESA} · Processo de pagamento CEABS · {ROTULO}')
    c.drawRightString(larg - 31, 20, f'Página {i + 1} de {total}')
    c.save()
    pg.merge_page(PdfReader(io.BytesIO(buf.getvalue())).pages[0])
w.add_metadata({'/Title': f'Processo de pagamento CEABS - {ROTULO}', '/Author': EMPRESA, '/Creator': EMPRESA, '/Producer': EMPRESA})
PROCESSO = os.path.join(SAIDA, f'PROCESSO PAGAMENTO CEABS - {SUFIXO}.pdf')
with open(PROCESSO, 'wb') as f:
    w.write(f)

# ---------- e-mail para a diretoria ----------
TH = 'background:#1F3A5F;color:#fff;padding:4px 8px;border:1px solid #9BB7D4;font-size:11px;text-align:left'
TD = 'padding:3px 8px;border:1px solid #D9D9D9;font-size:11px'
TT = 'padding:4px 8px;border:1px solid #9BB7D4;font-size:11px;background:#DDEBF7;font-weight:bold'
linhas_email = ''.join(f'<tr><td style="{TD}">{e(c)}</td><td style="{TD}">{e(st)}</td><td style="{TD};text-align:right">{n}</td></tr>'
                       for c, st, n in lin_cli[:-1])
email = f"""<html><head><meta charset="utf-8"><title>E-mail diretoria - CEABS {ROTULO}</title></head>
<body style="font-family:Calibri,Arial,sans-serif;font-size:11pt">
<p>Bom dia, Sr. Bruno Filho e Bruno Ribeiro.</p>
<p>Segue abaixo processo para autorização de pagamento referente à CEABS (rastreadores), medição de {MESES[NMES - 1]}/{ANO},
NFS-e {nfs_txt}, valor R$ {brl(total_pagar)[3:]}, vencimento {vencimento}, conforme relatório abaixo e processo em anexo
(FORPAG, notas fiscais e boletos, relatório resumido e consolidado).</p>
<p>Foram cobrados {n_equip} equipamentos (mês anterior: {n_equip_ant}), valor líquido R$ {brl(total_pagar)[3:]} (mês anterior: R$ {brl(liq_ant)[3:]}),
valor unitário de R$ {brl(total_pagar / n_equip)[3:]}. Dos veículos cobrados, {len(solicitado)} já tiveram o cancelamento solicitado
e {len(pedir)} serão solicitados agora ({', '.join(pedir)}).</p>
<p>Solicito autorização para pagamento.</p>
<table cellspacing="0" cellpadding="0" style="border-collapse:collapse;font-family:Calibri,Arial,sans-serif">
<tr><td colspan="3" style="{TH}">Medição {ROTULO} · NFS-e {nfs_txt} · vencimento {vencimento}</td></tr>
<tr><td style="{TH}">CLIENTE NO PERÍODO</td><td style="{TH}">STATUS</td><td style="{TH}">VEÍCULOS</td></tr>
{linhas_email}
<tr><td style="{TT}">Total geral</td><td style="{TT}"></td><td style="{TT};text-align:right">{lin_cli[-1][2]}</td></tr>
</table>
<p></p></body></html>"""
EMAIL = os.path.join(SAIDA, f'EMAIL DIRETORIA - CEABS {SUFIXO}.html')
open(EMAIL, 'w', encoding='utf-8').write(email)
shutil.rmtree(tmp, ignore_errors=True)
print('páginas:', total, '| relatório a partir da', inicio_rel + 1)
for p in (FORPAG_XLSX, FORPAG_PDF, PROCESSO, EMAIL):
    print('ok:', p)
