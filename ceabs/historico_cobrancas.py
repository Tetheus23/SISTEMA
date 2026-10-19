"""Planilha limpa da CEABS: histórico de cobranças item a item + painel dos casos e dos últimos 5 meses.

Uso: python3 -I historico_cobrancas.py <pasta_ceabs> <saida.xlsx> [planilha_revisada.xlsx]   (rodar depois de pipeline.sh)
Abas:
  Painel     resumo por status, últimos 5 meses (equipamentos, valor cheio, líquido, pagamento),
             notas desses meses, casos para tratar e pendências
  Histórico  uma linha por item da medição, com as colunas da planilha da CEABS e, no fim,
             mês apurado, vencimento, notas, valor sem imposto, pagamento, cliente, status e observações
VALOR SEM IMPOSTO: o líquido das notas (boleto) do mês e produto rateado entre os itens, proporcional ao valor a
faturar; a soma de cada grupo bate com o boleto ao centavo.
Planilha revisada: a mesma planilha devolvida com ajustes na aba Histórico. STATUS, CLIENTE NO PERÍODO, CLIENTE ATUAL,
DATA DE PAGAMENTO e OBSERVAÇÕES vêm dela; STATUS e CLIENTE ATUAL valem pela linha mais recente de cada veículo.
"""
import csv
import datetime as dt
import os
import re
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

BASE = sys.argv[1]
SAIDA = sys.argv[2]
REVISADA = sys.argv[3] if len(sys.argv) > 3 else None
EST = os.path.join(BASE, 'estrutura')
HOJE = dt.date(2026, 10, 19)
GRILLO = 'LOCADORA GRILLO E RIBEIRO LTDA'
CNPJ = '06025418000130'
RETENCAO = 0.0565      # IRRF 1% + CSRF 4,65% na NFS-e de mensalidade
AZUL, LARANJA, CINZA, BRANCO = '1F3864', 'C55A11', '595959', 'FFFFFF'
BRL = '"R$" #,##0.00'
DATA, MES = 'dd/mm/yyyy', 'mm/yyyy'
FINO = Side(style='thin', color='BFBFBF')
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']
STATUS = ['MANTER', 'PEDIR CANCELAMENTO', 'JÁ SOLICITADO CANCELAMENTO', 'CANCELADO E RETIRADO',
          'CANCELADO E RETIRADA A CONFIRMAR']
COR = dict(zip(STATUS, ['C6EFCE', 'F4B6B6', 'FFE699', 'D9D9D9', 'F8CBAD']))
ACAO = dict(zip(STATUS, ['Em uso. Nada a fazer.',
                         'Pedir à CEABS o cancelamento e a retirada.',
                         'Confirmar o aceite digital da rescisão e agendar a retirada (30 dias).',
                         'Encerrado.',
                         'Confirmar a retirada: a CEABS cobra R$ 400 por equipamento não devolvido.']))
CLASSE = {'Principal': 'Mensalidade', 'Taxa de Rescisão': 'Taxa de rescisão', 'Visita Improdutiva': 'Visita improdutiva'}
ORDEM_TIPO = {'Principal': 0, 'Taxa de Rescisão': 1, 'Visita Improdutiva': 2}
L0 = 2  # aba Histórico: linha 1 = totais do que estiver filtrado, linha 2 = cabeçalho


def ler(nome):
    return list(csv.DictReader(open(os.path.join(EST, nome), encoding='utf-8')))


def iso(s):
    s = (s or '')[:10]
    return dt.date.fromisoformat(s) if re.match(r'\d{4}-\d{2}-\d{2}$', s) else None


def brd(s):
    m = re.match(r'(\d{2})/(\d{2})/(\d{4})', s or '')
    return dt.date(int(m[3]), int(m[2]), int(m[1])) if m else None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def num_br(s):
    return float(s.replace('.', '').replace(',', '.'))


def mes_data(m):
    return dt.date(int(m[:4]), int(m[5:7]), 1)


def mes_txt(m):
    return f'{MESES[int(m[5:7]) - 1]}/{m[:4]}'


def produto_doc(doc):
    return 'Monitoramento GSM' if doc.startswith('Fatura') or '(monitoramento)' in doc else 'Telemetria'


def nome_doc(doc):
    return re.sub(r'\s*\((monitoramento|gestão de frotas)\)', '', doc).replace(' no controle', '')


def rateia(total, pesos):
    """Divide o total proporcionalmente aos pesos, em centavos; a sobra do arredondamento vai para o maior peso."""
    soma = sum(pesos)
    partes = [round(total * p / soma, 2) if soma else round(total / len(pesos), 2) for p in pesos]
    if partes:
        partes[max(range(len(pesos)), key=lambda i: pesos[i])] += round(total - sum(partes), 2)
    return [round(p, 2) for p in partes]


# ---------- dados ----------
itens = ler('itens_cobrados.csv')
base = ler('base.csv')
faturas = ler('faturas.csv')
retiradas = ler('retiradas.csv')
equip = {e['chassi']: e for e in ler('equipamentos.csv')}

info = {}  # chassi -> placa, modelo e situação atual na frota (base.csv já resolve os chassis digitados errado)
for r in ler('frota/veiculos_frota.csv'):
    info[r['ch']] = dict(placa=r['placa'], modelo=r['marca_modelo'], situacao=r['situacao_atual'], cliente=r['cliente_atual'])
local = {}
for b in sorted(base, key=lambda x: x['mes_uso']):
    info[b['chassi']] = dict(placa=b['placa'], modelo=b['modelo'], situacao=b['situacao_frota'], cliente=b['cliente_atual'])
    if b['tipo'] == 'Principal':
        local[(b['mes_uso'], b['chassi'])] = b['local_no_mes']


def dia(v):
    return v.date() if isinstance(v, dt.datetime) else v


def chave(codigo, mes, produto, valor, chassi):
    return codigo, mes, produto, round(valor or 0, 2), chassi


def cliente_no_mes(mes, ch, r=None):
    if r is not None and id(r) in revisado:
        return revisado[id(r)]['CLIENTE NO PERÍODO DA MEDIÇÃO'] or ''
    if (mes, ch) in local:
        return local[(mes, ch)]
    antes = [m for (m, c) in local if c == ch and m <= mes]  # taxa/visita sem mensalidade no mês: último local conhecido
    return local[(max(antes), ch)] if antes else ''


def cliente_atual(ch):
    if ch in cliente_rev:
        return cliente_rev[ch]
    f = info.get(ch)
    if not f:
        return 'Não localizado na frota'
    if f['situacao'] == 'vendido':
        return 'VENDIDO'
    return f['cliente'] or f['situacao'].replace('_', ' ').upper()


# ---------- notas: líquido de cada documento ----------
docs = defaultdict(list)  # (mês, classe) -> notas
estimado = set()
for f in faturas:
    if 'avulsa' in f['tipo_cobranca'] or f['tipo_cobranca'].startswith('Equipamento'):
        continue  # fora da medição: aparecem só nas pendências do painel
    docs[(f['mes_uso'], f['tipo_cobranca'])].append(f)
for grupo in docs.values():
    for f in grupo:
        f['bruto'] = float(f['valor_bruto'])
        if f['valor_liquido']:
            f['liq'] = float(f['valor_liquido'])
            continue
        junto = re.search(r'Pago junto com a (\d+): R\$ ([\d.,]+)', f['observacao'])
        par = junto and next((g for g in grupo if re.search(rf'\b{junto[1]}\b', g['documento'])), None)
        if par:
            f['liq'] = round(num_br(junto[2]) * f['bruto'] / (f['bruto'] + float(par['valor_bruto'])), 2)
        else:
            reter = f['documento'].startswith('NFS-e') and f['tipo_cobranca'] == 'Mensalidade'
            f['liq'] = round(f['bruto'] * (1 - RETENCAO), 2) if reter else f['bruto']
            estimado.add(f['mes_uso'])

# ---------- grupos mês x produto e rateio ----------
por_classe = defaultdict(list)
for r in itens:
    por_classe[(r['mes_uso'], CLASSE[r['tipo']])].append(r)
fiscal = {}  # id(item) -> dados da nota
for (mes, classe), rs in por_classe.items():
    ds = docs.get((mes, classe), [])
    if not ds:
        continue
    grupos = [(rs, ds)]
    if classe == 'Mensalidade':
        prods = {r['produto'] for r in rs}
        separado = [([r for r in rs if r['produto'] == p], [d for d in ds if produto_doc(d['documento']) == p]) for p in prods]
        bate = all(dd and abs(sum(num(r['valor']) or 0 for r in ri) - sum(d['bruto'] for d in dd)) <= max(1, 0.02 * sum(d['bruto'] for d in dd))
                   for ri, dd in separado) and sum(len(dd) for _, dd in separado) == len(ds)
        if bate:
            grupos = separado
    for ri, dd in grupos:
        vencs = sorted(brd(d['vencimento']) for d in dd if brd(d['vencimento']))
        pagos = [brd(d['pago_em']) for d in dd]
        dados = dict(notas=' + '.join(nome_doc(d['documento']) for d in dd),
                     venc=vencs[0] if vencs else None,
                     pago=max(pagos) if pagos and all(pagos) else None)
        for r, v in zip(ri, rateia(round(sum(d['liq'] for d in dd), 2), [num(r['valor']) or 0 for r in ri])):
            fiscal[id(r)] = dict(dados, liquido=v)

# ---------- status por veículo ----------
status, pedido = {}, {}
for r in retiradas:
    ch, g = r['chassi'], r['grupo'][0]
    if g == '1':
        status[ch] = 'JÁ SOLICITADO CANCELAMENTO' if r['cancelamento'] else 'PEDIR CANCELAMENTO'
        pedido[ch] = r['cancelamento']
    elif g == '3':
        status[ch] = 'JÁ SOLICITADO CANCELAMENTO'
        pedido[ch] = (re.search(r'\d{2}/\d{2}/\d{4}', r['situacao']) or [''])[0]
    else:
        status[ch] = 'CANCELADO E RETIRADO' if r['retirada'] else 'CANCELADO E RETIRADA A CONFIRMAR'
ult_mes = max(r['mes_uso'] for r in itens)
atuais = {r['chassi'] for r in itens if r['mes_uso'] == ult_mes and r['tipo'] == 'Principal'}
for r in itens:
    status.setdefault(r['chassi'], 'MANTER' if r['chassi'] in atuais else 'CANCELADO E RETIRADO')
sem_item = set(status) - {r['chassi'] for r in itens}
if sem_item:
    print('AVISO: chassis da lista de retiradas sem item na medição:', sorted(sem_item))

# ---------- ajustes da planilha revisada ----------
revisado, cliente_rev = {}, {}  # id(item) -> linha revisada; chassi -> cliente atual revisado
if REVISADA:
    ws_r = load_workbook(REVISADA, data_only=True)['Histórico']
    cab = [c.value for c in ws_r[L0]]
    fila = defaultdict(list)
    for x in ws_r.iter_rows(min_row=L0 + 1, values_only=True):
        r = dict(zip(cab, x))
        if r.get('Código'):
            mes = dia(r['MÊS/ANO APURADO'])
            fila[chave(r['Código'], mes.strftime('%Y-%m') if mes else '', r['Produto'], r['Valor a Faturar'], r['Chassi'])].append(r)
    for r in itens:
        lst = fila.get(chave(r['codigo'], r['mes_uso'], r['produto'], num(r['valor']), r['chassi']))
        if lst:
            revisado[id(r)] = lst.pop(0)
    sobra = sum(len(v) for v in fila.values())
    print(f'revisada: {len(revisado)} de {len(itens)} itens casados', f'| AVISO: {sobra} linhas da revisada sem item' if sobra else '')
    # situação atual = a da linha mais recente do veículo (mensalidade antes de taxa/visita)
    for r in sorted((r for r in itens if id(r) in revisado), key=lambda x: (x['mes_uso'], -ORDEM_TIPO[x['tipo']])):
        rv = revisado[id(r)]
        if rv['STATUS'] in STATUS:
            status[r['chassi']] = rv['STATUS']
        if rv['CLIENTE ATUAL']:
            cliente_rev[r['chassi']] = rv['CLIENTE ATUAL']

# ---------- planilha ----------
wb = Workbook()
painel = wb.active
painel.title = 'Painel'
hist = wb.create_sheet('Histórico')

ORIGINAIS = ['Cliente', 'Proprietário', 'CPF/CNPJ', 'Código', 'Chassi', 'Placa', 'Produto', 'Monitoramento', 'Pró-Rata',
             'Valor a Faturar', 'Data de Instalação', 'Data de Cancelamento', 'Data de Retirada']
NOVAS = ['MÊS/ANO APURADO', 'MÊS/ANO VENCIMENTO', 'NOTAS FISCAIS', 'VALOR SEM IMPOSTO', 'VENCIMENTO', 'DATA DE PAGAMENTO',
         'CLIENTE NO PERÍODO DA MEDIÇÃO', 'CLIENTE ATUAL', 'STATUS', 'OBSERVAÇÕES']
TITULOS = ORIGINAIS + NOVAS
LARG = [12, 12, 15, 17, 20, 10, 18, 14, 11, 13, 12, 14, 12, 11, 13, 26, 13, 13, 13, 30, 28, 30, 30]
FORMATO = {8: BRL, 9: BRL, 10: BRL, 11: DATA, 12: DATA, 13: DATA, 14: MES, 15: MES, 17: BRL, 18: DATA, 19: DATA}

linhas = []
for r in sorted(itens, key=lambda x: (x['mes_uso'], -ORDEM_TIPO[x['tipo']], x['placa'] or (info.get(x['chassi']) or {}).get('placa', '')),
                reverse=True):
    fz = fiscal.get(id(r), {})
    rv = revisado.get(id(r), {})
    venc = fz.get('venc')
    placa = r['placa'] or (info.get(r['chassi']) or {}).get('placa', '')
    linhas.append([GRILLO, GRILLO, CNPJ, r['codigo'], r['chassi'], placa, r['produto'], num(r['mensal']), num(r['prorata']),
                   num(r['valor']), iso(r['instalacao']), iso(r['cancelamento']), iso(r['retirada']),
                   mes_data(r['mes_uso']), dt.date(venc.year, venc.month, 1) if venc else None, fz.get('notas', ''),
                   fz.get('liquido'), venc, dia(rv.get('DATA DE PAGAMENTO')) or fz.get('pago'),
                   cliente_no_mes(r['mes_uso'], r['chassi'], r), cliente_atual(r['chassi']), status[r['chassi']],
                   rv.get('OBSERVAÇÕES')])

for c, t in enumerate(TITULOS, 1):
    cel = hist.cell(L0, c, t)
    cel.font = Font(bold=True, color=BRANCO)
    cel.fill = PatternFill('solid', fgColor=AZUL if c <= len(ORIGINAIS) else LARANJA)
    cel.alignment = Alignment(wrap_text=True, vertical='center')
    hist.column_dimensions[get_column_letter(c)].width = LARG[c - 1]
hist.row_dimensions[L0].height = 46
for i, lin in enumerate(linhas, L0 + 1):
    for c, v in enumerate(lin, 1):
        cel = hist.cell(i, c, v)
        if c in FORMATO:
            cel.number_format = FORMATO[c]
ult = L0 + len(linhas)
col = {t: get_column_letter(i) for i, t in enumerate(TITULOS, 1)}
hist['A1'] = 'Total do que está filtrado:'
hist['A1'].font = Font(bold=True, color=AZUL)
for t, f in (('Código', 103), ('Valor a Faturar', 109), ('VALOR SEM IMPOSTO', 109)):
    cel = hist[f'{col[t]}1']
    cel.value = f'=SUBTOTAL({f},{col[t]}{L0 + 1}:{col[t]}{ult})'
    cel.font = Font(bold=True)
    cel.number_format = '0" itens"' if f == 103 else BRL
tab = Table(displayName='Historico', ref=f'A{L0}:{get_column_letter(len(TITULOS))}{ult}')
tab.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True)
hist.add_table(tab)
hist.freeze_panes = f'A{L0 + 1}'
faixa = f'{col["STATUS"]}{L0 + 1}:{col["STATUS"]}{ult}'
for s in STATUS:
    hist.conditional_formatting.add(faixa, CellIsRule(operator='equal', formula=[f'"{s}"'], fill=PatternFill('solid', fgColor=COR[s])))
dv = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True)
hist.add_data_validation(dv)
dv.add(faixa)


# ---------- painel ----------
def titulo_secao(ws, linha, texto, nota=None):
    ws.cell(linha, 1, texto).font = Font(bold=True, size=13, color=AZUL)
    if nota:
        ws.cell(linha + 1, 1, nota).font = Font(italic=True, color=CINZA)


def tabela(ws, linha, titulos, rows, formatos=None, cor_status=None):
    for c, t in enumerate(titulos, 1):
        cel = ws.cell(linha, c, t)
        cel.font = Font(bold=True, color=BRANCO)
        cel.fill = PatternFill('solid', fgColor=AZUL)
        cel.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[linha].height = 46
    for i, lin in enumerate(rows, linha + 1):
        for c, v in enumerate(lin, 1):
            cel = ws.cell(i, c, v)
            cel.border = Border(bottom=FINO)
            cel.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str) and len(v) > 28,
                                      horizontal='center' if isinstance(v, dt.date) else None)
            if formatos and c in formatos:
                cel.number_format = formatos[c]
        if cor_status:
            st = lin[cor_status - 1]
            ws.cell(i, cor_status).fill = PatternFill('solid', fgColor=COR.get(st, BRANCO))
    return linha + len(rows) + 1


ws = painel
for c, w in enumerate([30, 13, 38, 20, 26, 16, 16, 18, 16, 40], 1):
    ws.column_dimensions[get_column_letter(c)].width = w
ws['A1'] = 'CEABS (rastreadores) - histórico de cobranças e situação dos equipamentos'
ws['A1'].font = Font(bold=True, size=16, color=AZUL)
ws['A2'] = (f'Contrato 40010217 - {GRILLO} - posição em {HOJE:%d/%m/%Y}. Fonte: planilhas de medição da CEABS, '
            'NFS-e/faturas e boletos, base da frota. O detalhe item a item está na aba Histórico.')
ws['A2'].font = Font(italic=True, color=CINZA)
if REVISADA:
    ws['A3'] = f'Status e clientes revisados pelo Matheus em {HOJE:%d/%m/%Y} (principalmente o mês apurado {mes_txt(ult_mes)}).'
    ws['A3'].font = Font(italic=True, bold=True, color=LARANJA)

# 1. resumo por status
veic_ult = defaultdict(lambda: [set(), 0.0, 0.0])
for lin in linhas:
    if lin[13] == mes_data(ult_mes):
        v = veic_ult[lin[21]]
        v[1] += lin[9] or 0
        v[2] += lin[16] or 0
n_status = Counter(status.values())
titulo_secao(ws, 4, '1. Situação dos veículos com rastreador CEABS',
             f'Cada veículo aparece uma vez, pela situação atual. Valores = cobrança da medição de {mes_txt(ult_mes)}.')
linha = tabela(ws, 6, ['Status', 'Veículos', 'O que fazer', f'Cobrado em {mes_txt(ult_mes)} (cheio)',
                       f'Cobrado em {mes_txt(ult_mes)} (líquido)'],
               [[s, n_status[s], ACAO[s], round(veic_ult[s][1], 2), round(veic_ult[s][2], 2)] for s in STATUS],
               {4: BRL, 5: BRL}, cor_status=1)
ws.cell(linha, 1, 'Total').font = Font(bold=True)
ws.cell(linha, 2, f'=SUM(B7:B{linha - 1})').font = Font(bold=True)
for c in (4, 5):
    cel = ws.cell(linha, c, f'=SUM({get_column_letter(c)}7:{get_column_letter(c)}{linha - 1})')
    cel.font, cel.number_format = Font(bold=True), BRL

# 2. últimos 5 meses
meses5 = sorted({r['mes_uso'] for r in itens})[-5:]
docs5 = [f for f in faturas if f['mes_uso'] in meses5 and not f['tipo_cobranca'].startswith('Equipamento')]
resumo = []
for m in meses5:
    rs = [lin for lin in linhas if lin[13] == mes_data(m)]
    principal = [lin for lin in rs if lin[6] not in ('Taxa de Rescisão', 'Visita Improdutiva')]
    dm = [f for f in docs5 if f['mes_uso'] == m]
    vencs = sorted({brd(f['vencimento']) for f in dm})
    pagos = [brd(f['pago_em']) for f in dm]
    resumo.append([mes_data(m), vencs[0] if vencs else None, ' + '.join(nome_doc(f['documento']) for f in dm),
                   len({lin[3] for lin in principal}), round(sum(lin[9] or 0 for lin in principal), 2),
                   round(sum(lin[9] or 0 for lin in rs if lin not in principal), 2), round(sum(lin[9] or 0 for lin in rs), 2),
                   round(sum(lin[16] or 0 for lin in rs), 2), max(pagos) if pagos and all(pagos) else None])
L5 = linha + 3
titulo_secao(ws, L5, '2. Últimos 5 meses',
             'Valor cheio = soma da planilha da CEABS. Valor líquido = boletos (notas menos retenções). '
             'Pagamento em branco = sem registro: preencher.')
fim5 = tabela(ws, L5 + 2, ['Mês apurado', 'Vencimento', 'Notas fiscais', 'Equipamentos cobrados', 'Mensalidades',
                           'Taxas e visitas', 'Valor cheio', 'Valor líquido (boletos)', 'Data de pagamento'],
              resumo, {1: MES, 2: DATA, 5: BRL, 6: BRL, 7: BRL, 8: BRL, 9: DATA})
graf = BarChart()
graf.type, graf.grouping = 'col', 'clustered'
graf.title = 'Valor cheio x líquido por mês'
graf.y_axis.numFmt = '#,##0'
graf.y_axis.title = 'R$'
graf.add_data(Reference(ws, min_col=7, max_col=8, min_row=L5 + 2, max_row=fim5 - 1), titles_from_data=True)
graf.set_categories(Reference(ws, min_col=1, min_row=L5 + 3, max_row=fim5 - 1))
graf.height, graf.width = 7.5, 16
ws.add_chart(graf, f'K{L5}')

# 3. notas dos últimos 5 meses
L6 = fim5 + 2
titulo_secao(ws, L6, '3. Notas fiscais desses meses', 'Uma linha por nota. Onde não há pagamento registrado, a data fica em branco.')
fim6 = tabela(ws, L6 + 2, ['Mês apurado', 'Nota', 'Tipo', 'Vencimento', 'Valor bruto', 'Valor líquido', 'Data de pagamento',
                           'Situação'],
              [[mes_data(f['mes_uso']), nome_doc(f['documento']), f['tipo_cobranca'], brd(f['vencimento']), float(f['valor_bruto']),
                next((g['liq'] for g in docs.get((f['mes_uso'], f['tipo_cobranca']), []) if g['documento'] == f['documento']),
                     num(f['valor_liquido'])), brd(f['pago_em']), f['situacao']]
               for f in sorted(docs5, key=lambda x: (x['mes_uso'], x['documento']), reverse=True)],
              {1: MES, 4: DATA, 5: BRL, 6: BRL, 7: DATA})

# 4. casos para tratar
ultimo = {}
for r in itens:
    ultimo[r['chassi']] = max(ultimo.get(r['chassi'], ''), r['mes_uso'])
sinistrado = {r['chassi'] for r in retiradas if r['situacao'] == 'Sinistrado'}


def acao_caso(ch, st):
    if st == 'PEDIR CANCELAMENTO':
        return ('Sinistrado: confirmar se voltou à operação; se não, pedir cancelamento e retirada' if ch in sinistrado
                else 'Pedir cancelamento e retirada')
    if st == 'JÁ SOLICITADO CANCELAMENTO':
        return f'Pedido em {pedido.get(ch) or "set/2026"}: confirmar o aceite da rescisão e agendar a retirada'
    e = equip.get(ch)
    if not e:
        return 'Cancelado: confirmar se o equipamento foi retirado'
    if e['acao'].startswith('CONTESTAR'):
        return 'Na cobrança de R$ 400: contestar (retirada pedida e taxa paga antes da venda)'
    if e['situacao_frota'] == 'vendido':
        return f'Na cobrança de R$ 400: vendido ({e["comprador"] or "comprador"}), ver a retirada com o comprador'
    return 'Na cobrança de R$ 400: ainda é nosso, agendar a retirada'


casos = sorted(((st, (info.get(ch) or {}).get('placa') or ch, (info.get(ch) or {}).get('modelo', ''), ch, cliente_atual(ch),
                 mes_data(ultimo[ch]), acao_caso(ch, st))
                for ch, st in status.items() if st in STATUS[1:3] + STATUS[4:]), key=lambda x: (STATUS.index(x[0]), x[1]))
L7 = fim6 + 2
titulo_secao(ws, L7, '4. Casos para tratar', f'{len(casos)} veículos. Os demais estão em uso (MANTER) ou já cancelados e retirados.')
fim7 = tabela(ws, L7 + 2, ['Status', 'Placa', 'Modelo', 'Chassi', 'Cliente atual', 'Último mês cobrado', 'O que fazer'],
              [list(c) for c in casos], {6: MES}, cor_status=1)
for i in range(L7 + 3, fim7):  # "O que fazer" é longo: ocupa até a coluna J
    ws.merge_cells(start_row=i, start_column=7, end_row=i, end_column=10)
ws.merge_cells(start_row=L7 + 2, start_column=7, end_row=L7 + 2, end_column=10)

# 5. pendências e como ler
L8 = fim7 + 2
titulo_secao(ws, L8, '5. Pendências fora da medição')
pend = ['Vencimento da fatura de ago/2026 (NFS-e 10276/10277) em 25/09/2026, fora do padrão (dia 3 do 2º mês após o uso '
        '= 03/10). As duas estão em aberto e vencidas.',
        'NFS-e 8978 (visitas) e 8979 (taxas de rescisão) de jul/2026 e NFS-e 2324 (visitas de jan/2026): contestadas com a CEABS.',
        'Cobrança de equipamentos não devolvidos (aviso de 13/10/2026): R$ 400 por equipamento a partir de 20/10. '
        'São os veículos "Na cobrança de R$ 400" da seção 4.',
        'Fora das medições: NFS-e 124199 (equipamentos, R$ 800, venc. 13/10/2025, contestada) e NFS-e 32928 (visita avulsa emitida em ago/2025, R$ 45).']
for i, p in enumerate(pend, L8 + 1):
    ws.cell(i, 1, f'- {p}')
L9 = L8 + len(pend) + 2
titulo_secao(ws, L9, 'Como ler a aba Histórico')
notas = ['Colunas em azul = planilha da CEABS. Colunas em laranja = conferência (incluídas no fim). Linha 1 soma só o que estiver filtrado.',
         'VALOR SEM IMPOSTO = boleto (líquido das notas) do mês e produto dividido entre os itens, proporcional ao valor de cada um '
         '(pró-rata recebe proporcionalmente menos). A soma de cada mês bate com os boletos.',
         'STATUS = situação atual do veículo, repetida em todas as linhas dele. Placa em branco na CEABS foi preenchida pela base da frota.',
         'Líquido calculado pela retenção padrão (5,65% na NFS-e de mensalidade) onde não há boleto registrado: '
         + ', '.join(mes_txt(m) for m in sorted(estimado)) + '. Ago-out/2023 não têm nota no histórico.']
for i, n in enumerate(notas, L9 + 1):
    ws.cell(i, 1, f'- {n}').font = Font(color=CINZA)
ws.freeze_panes = 'A4'
ws.sheet_view.showGridLines = False
for w in (ws, hist):
    w.page_setup.orientation = 'landscape'
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
wb.calculation.fullCalcOnLoad = True
wb.save(SAIDA)

# ---------- conferência ----------
print('itens:', len(linhas), '| status:', dict(n_status))
dif = []
for (mes, classe), ds in sorted(docs.items()):
    rs = por_classe.get((mes, classe), [])
    if rs:
        soma = round(sum(fiscal[id(r)]['liquido'] for r in rs if id(r) in fiscal), 2)
        liq = round(sum(d['liq'] for d in ds), 2)
        if abs(soma - liq) > 0.005:
            dif.append((mes, classe, soma, liq))
print('grupos com rateio diferente do boleto:', dif or 'nenhum')
print('itens sem nota:', Counter(r['mes_uso'] for r in itens if id(r) not in fiscal))
print('notas sem item:', sorted(k for k in docs if k not in por_classe))
for r in resumo:
    print(r[0].strftime('%m/%Y'), r[3], 'equip | cheio', r[6], '| líquido', r[7], '| pago', r[8])
print('ok:', SAIDA)
