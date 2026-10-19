"""Monta a planilha final da auditoria CEABS (Excel) a partir de estrutura/*.csv.

Uso: python3 -I montar_painel.py <pasta_ceabs> <saida.xlsx>
Depois rodar dinamicas.py (LibreOffice) para inserir as tabelas dinâmicas reais.
Abas: Painel, (Dinâmicas), Ago-26 por veículo, Contestar CEABS, Equipamentos, Retiradas, Faturas, Base, Como ler.
"""
import csv
import datetime as dt
import os
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = sys.argv[1]
SAIDA = sys.argv[2]
EST = os.path.join(BASE, 'estrutura')
HOJE = dt.date(2026, 10, 19)
MES_ATUAL = '2026-08'
ROTULO = {'OK': 'OK', 'SEM USO': 'CANCELAR E RETIRAR', 'CONTESTAR': 'CONTESTAR', 'VERIFICAR': 'VERIFICAR'}
COR = {'OK': 'C6EFCE', 'CANCELAR E RETIRAR': 'FFE699', 'CONTESTAR': 'F4B6B6', 'VERIFICAR': 'D9D9D9'}
SITUACAO = {'sinistro': 'Sinistrado', 'disponivel_venda': 'À venda', 'vendido': 'Vendido', 'leilao': 'Leilão',
            'interno': 'Uso interno sem uso registrado'}
AZUL, CINZA, BRANCO = '1F3864', '595959', 'FFFFFF'
BRL = '"R$" #,##0.00'
FINO = Side(style='thin', color='BFBFBF')
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']


def ler(nome):
    return list(csv.DictReader(open(os.path.join(EST, nome), encoding='utf-8')))


def mes_txt(m):
    return f'{MESES[int(m[5:7]) - 1]}/{m[2:4]}'


def brl(v):
    return f'R$ {v:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def cabecalho(ws, linha, titulos, larguras=None):
    for i, t in enumerate(titulos, 1):
        c = ws.cell(linha, i, t)
        c.font = Font(bold=True, color=BRANCO)
        c.fill = PatternFill('solid', fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical='center')
        c.border = Border(bottom=FINO)
    ws.row_dimensions[linha].height = 32
    for i, w in enumerate(larguras or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def tabela(ws, linha0, titulos, linhas, larguras, moedas=(), titulo=None, nota=None):
    """Escreve título, nota, cabeçalho e linhas; aplica filtro, congela e formata moeda."""
    if titulo:
        ws.cell(1, 1, titulo).font = Font(bold=True, size=14, color=AZUL)
    if nota:
        ws.cell(2, 1, nota).font = Font(italic=True, color=CINZA)
    cabecalho(ws, linha0, titulos, larguras)
    for r, lin in enumerate(linhas, linha0 + 1):
        for c, v in enumerate(lin, 1):
            cel = ws.cell(r, c, v)
            cel.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str) and len(v) > 40)
            if c in moedas:
                cel.number_format = BRL
    ult = linha0 + max(len(linhas), 1)
    ws.auto_filter.ref = f'A{linha0}:{get_column_letter(len(titulos))}{ult}'
    ws.freeze_panes = ws.cell(linha0 + 1, 3)
    return ult


def pinta_status(ws, col, r1, r2):
    for rot, cor in COR.items():
        ws.conditional_formatting.add(f'{col}{r1}:{col}{r2}', CellIsRule(operator='equal', formula=[f'"{rot}"'],
                                                                          fill=PatternFill('solid', fgColor=cor)))


# ---------- dados ----------
base = ler('base.csv')
fat = ler('faturas.csv')
venc_mes = {}
for f in fat:
    if f['tipo_cobranca'] == 'Mensalidade' and f['vencimento'] and f['mes_uso'] not in venc_mes:
        dd, mm, aa = f['vencimento'].split('/')
        venc_mes[f['mes_uso']] = (f'{aa}-{mm}', f['vencimento'])
for b in base:
    if b['status'] == 'SEM USO':
        if b['mes_uso'] < MES_ATUAL:
            # meses passados: o relatório só aponta contestação; o "sem uso" vale para o mês atual (cancelar e retirar)
            b['status'], b['motivo'] = 'OK', 'Mensalidade do mês'
            if b['local_no_mes'] in ('SINISTRO', 'VENDIDO'):
                b['local_no_mes'] = ''
        else:
            sit = SITUACAO.get(b['categoria_local'], 'Sem cliente')
            ped = f'; cancelamento pedido em {b["pedido_cancelamento"]}' if b['pedido_cancelamento'] else ''
            b['motivo'] = f'{sit}: solicitar cancelamento e retirada{ped}'
    b['status_txt'] = ROTULO[b['status']]
    b['mes_venc'], b['venc_data'] = venc_mes.get(b['mes_uso'], ('', ''))
meses = sorted({b['mes_uso'] for b in base})

wb = Workbook()
# ---------- Base ----------
wsB = wb.active
wsB.title = 'Base'
colsB = [('Mês de uso', 'mes_uso', 10), ('Mês de vencimento', 'mes_venc', 11), ('Vencimento', 'venc_data', 11),
         ('Referência CEABS', 'ref_ceabs', 15), ('Tipo de cobrança', 'tipo', 16), ('Placa', 'placa', 10),
         ('Chassi', 'chassi', 19), ('Código CEABS', 'codigo_ceabs', 15), ('Modelo', 'modelo', 26),
         ('Valor cobrado', 'valor', 12), ('Status', 'status_txt', 18), ('Motivo', 'motivo', 60),
         ('Local / cliente no mês', 'local_no_mes', 28), ('Valor a contestar', 'valor_contestar', 12),
         ('Pedido de cancelamento', 'pedido_cancelamento', 12), ('Cancelamento CEABS', 'cancelamento_ceabs', 12),
         ('Instalação', 'instalacao', 11), ('Situação hoje (frota)', 'situacao_frota', 20),
         ('Cliente hoje (frota)', 'cliente_atual', 26), ('Evidência', 'evidencia', 60), ('Planilha de origem', 'planilha_origem', 30)]
linhasB = [[num(b[k]) if k in ('valor', 'valor_contestar') else b[k] for _, k, _ in colsB]
           for b in sorted(base, key=lambda x: (x['mes_uso'], x['status'], x['placa']))]
ultB = tabela(wsB, 1, [c[0] for c in colsB], linhasB, [c[2] for c in colsB], moedas=(10, 14))
wsB.freeze_panes = 'G2'
pinta_status(wsB, 'K', 2, ultB)
N = ultB  # última linha de dados da Base


def soma(crit):
    """SUMIFS sobre a Base (coluna J = valor cobrado)."""
    partes = ''.join(f',Base!${c}$2:${c}${N},"{v}"' for c, v in crit)
    return f'=SUMIFS(Base!$J$2:$J${N}{partes})'


def conta(crit):
    partes = ','.join(f'Base!${c}$2:${c}${N},"{v}"' for c, v in crit)
    return f'=COUNTIFS({partes})'


# ---------- Painel ----------
ws = wb.create_sheet('Painel', 0)
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 2
for c in 'BCDEFGH':
    ws.column_dimensions[c].width = 17
ws.column_dimensions['I'].width = 2
ws['B1'] = 'CEABS - Auditoria dos rastreadores'
ws['B1'].font = Font(bold=True, size=18, color=AZUL)
ws['B2'] = (f'Medição de ago/26 (uso em agosto, vencimento 25/09/2026) e histórico de {mes_txt(meses[0])} a {mes_txt(meses[-1])}. '
            f'Fontes: 736 e-mails CEABS, medições, NFS-e, base da frota, sinistros, diesel, Ticketlog e rastreador. Gerado em {HOJE:%d/%m/%Y}.')
ws['B2'].font = Font(italic=True, color=CINZA)
ws.merge_cells('B2:H2')
ws['B2'].alignment = Alignment(wrap_text=True)
ws.row_dimensions[2].height = 30

atual = [b for b in base if b['mes_uso'] == MES_ATUAL]
eq = ler('equipamentos.csv')
eq_cont = sum(float(e['valor']) for e in eq if e['acao'].startswith('CONTESTAR') or 'não pagar' in e['acao'])
ret = ler('retiradas.csv')
pend = [r for r in ret if r['grupo'][0] in '23' and not r['situacao'].startswith('Retirado')]
cards = [
    ('B', 4, f'Medição {mes_txt(MES_ATUAL)}', soma([('A', MES_ATUAL)]), f'{len([b for b in atual if b["tipo"] == "Principal"])} rastreadores cobrados', 'DDEBF7'),
    ('D', 4, 'Em uso (OK)', soma([('A', MES_ATUAL), ('K', ROTULO['OK'])]), conta([('A', MES_ATUAL), ('K', ROTULO['OK'])]), COR['OK']),
    ('F', 4, 'Cancelar e retirar', soma([('A', MES_ATUAL), ('K', ROTULO['SEM USO'])]), conta([('A', MES_ATUAL), ('K', ROTULO['SEM USO'])]), COR['CANCELAR E RETIRAR']),
    ('B', 8, 'Contestar com a CEABS', f'=SUM(Base!$N$2:$N${N})', 'cobranças após pedido, taxas e duplicidades', COR['CONTESTAR']),
    ('D', 8, 'Retiradas pendentes', len(pend), 'cancelados sem retirada confirmada', 'FCE4D6'),
    ('F', 8, 'Equipamentos (vence 20/10)', 4800, f'{brl(eq_cont)} contestáveis', 'F4B6B6'),
]
for col, lin, tit, val, sub, cor in cards:
    c2 = chr(ord(col) + 1)
    for r in range(lin, lin + 3):
        for cc in (col, c2):
            ws[f'{cc}{r}'].fill = PatternFill('solid', fgColor=cor)
    ws[f'{col}{lin}'] = tit
    ws[f'{col}{lin}'].font = Font(bold=True, color=AZUL)
    ws[f'{col}{lin + 1}'] = val
    ws[f'{col}{lin + 1}'].font = Font(bold=True, size=16)
    ws[f'{col}{lin + 1}'].number_format = BRL if not isinstance(val, int) or val > 1000 else '0'
    ws[f'{col}{lin + 2}'] = sub
    ws[f'{col}{lin + 2}'].font = Font(color=CINZA, size=9)
    if isinstance(sub, str) and sub.startswith('=COUNTIFS'):
        ws[f'{col}{lin + 2}'].number_format = '0" rastreadores"'

# ações
ra = 12
ws[f'B{ra}'] = 'O que fazer agora'
ws[f'B{ra}'].font = Font(bold=True, size=12, color=AZUL)
sem_uso_atual = [b for b in atual if b['status'] == 'SEM USO']
ACOES_AQUI = True

# resumo mês x status (alimenta o gráfico)
r0 = 22
ws[f'B{r0 - 1}'] = 'Valor cobrado por mês de uso e status'
ws[f'B{r0 - 1}'].font = Font(bold=True, size=12, color=AZUL)
cab = ['Mês de uso', 'Vencimento', 'OK', 'CANCELAR E RETIRAR', 'CONTESTAR', 'VERIFICAR', 'Total']
for i, t in enumerate(cab):
    c = ws.cell(r0, 2 + i, t)
    c.font = Font(bold=True, color=BRANCO)
    c.fill = PatternFill('solid', fgColor=AZUL)
    c.alignment = Alignment(wrap_text=True, horizontal='center')
for i, m in enumerate(meses, 1):
    r = r0 + i
    ws.cell(r, 2, mes_txt(m))
    ws.cell(r, 3, venc_mes.get(m, ('', ''))[1])
    for j, st in enumerate(['OK', 'CANCELAR E RETIRAR', 'CONTESTAR', 'VERIFICAR']):
        c = ws.cell(r, 4 + j, soma([('A', m), ('K', st)]))
        c.number_format = BRL + ';-' + BRL + ';"-"'
    c = ws.cell(r, 8, f'=SUM(D{r}:G{r})')
    c.number_format = BRL
    c.font = Font(bold=True)
rf = r0 + len(meses)
graf = BarChart()
graf.type, graf.grouping, graf.overlap = 'col', 'stacked', 100
graf.title = 'Cobrança mensal por status'
graf.y_axis.title = 'R$'
graf.y_axis.numFmt = '#,##0'
graf.y_axis.scaling.min = 0
graf.y_axis.delete = False
graf.x_axis.delete = False
graf.legend.position = 'b'
graf.add_data(Reference(ws, min_col=4, max_col=7, min_row=r0, max_row=rf), titles_from_data=True)
graf.set_categories(Reference(ws, min_col=2, min_row=r0 + 1, max_row=rf))
for s, cor in zip(graf.series, ('70AD47', 'FFC000', 'C00000', 'A6A6A6')):
    s.graphicalProperties.solidFill = cor
    s.graphicalProperties.line.solidFill = cor
graf.height, graf.width = 9, 24
ws.add_chart(graf, f'B{rf + 2}')

acoes = [
    ('Até 20/10 (amanhã)', f'Cobrança de 12 equipamentos não devolvidos (R$ 4.800). Agendar a retirada do que ainda é nosso e contestar os que '
                           f'tiveram retirada pedida e taxa paga ({brl(eq_cont)}). Detalhe na aba Equipamentos.'),
    ('Vencidas', 'NFS-e 10276/10277 (ago/26, R$ 4.099,69 líquido) vencidas em 25/09. A CEABS cancela o contrato com 60 dias de atraso '
                 '(24/11) e negativa de forma automática. NFS-e 8978/8979 (jul/26) vencidas em 03/09: contestação em aberto, 60 dias em 02/11.'),
    ('Retiradas', f'{len(sem_uso_atual)} rastreadores para solicitar cancelamento e retirada, {len(pend)} cancelados ou pedidos em set/26 sem '
                  'retirada confirmada. Sem o aceite digital da rescisão, o contrato segue ativo e cobrando. Lista na aba Retiradas; '
                  'e-mail para a equipe já redigido.'),
    ('Contestar', 'Pedir à CEABS a conferência das cobranças da aba Contestar CEABS (taxa do RMF6A73, visitas duplicadas, '
                  'veículos cobrados depois do pedido de retirada) e a revisão da NFS-e 8979 (e-mail já redigido).'),
    ('Fiscal', 'Corrigir o endereço do tomador nas próximas notas (Rua Nicarágua, 630). NFS-e de jan e fev/26 vieram sem retenção '
               'federal: conferir se a Grillo recolheu o DARF dos 5,65% que reteve no pagamento.'),
    ('Vencimento', 'A fatura de ago/26 veio com vencimento em 25/09. O padrão é o dia 3 do 2º mês após o uso (03/10). '
                   'Pedir à CEABS que mantenha o padrão.'),
]
for i, (t, txt) in enumerate(acoes, 1):
    ws.cell(ra + i, 2, t).font = Font(bold=True)
    ws.cell(ra + i, 2).alignment = Alignment(wrap_text=True, vertical='top')
    c = ws.cell(ra + i, 3, txt)
    ws.merge_cells(start_row=ra + i, start_column=3, end_row=ra + i, end_column=8)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.row_dimensions[ra + i].height = 46

# ---------- Ago-26 por veículo ----------
wsA = wb.create_sheet(f'{mes_txt(MES_ATUAL).capitalize().replace("/", "-")} por veículo', 1)
ordem = {'CANCELAR E RETIRAR': 0, 'CONTESTAR': 1, 'VERIFICAR': 2, 'OK': 3}
acao_ret = {r['chassi']: r['acao'] for r in ret if r['grupo'].startswith('1')}
linhasA = [[b['status_txt'], b['placa'], b['chassi'], b['modelo'], num(b['valor']), b['motivo'], b['local_no_mes'],
            b['pedido_cancelamento'], b['situacao_frota'], b['cliente_atual'], acao_ret.get(b['chassi'], '')]
           for b in sorted((b for b in atual if b['tipo'] == 'Principal'), key=lambda b: (ordem[b['status_txt']], b['placa']))]
ultA = tabela(wsA, 4, ['Status', 'Placa', 'Chassi', 'Modelo', 'Valor ago/26', 'Motivo', 'Local / cliente em agosto', 'Pedido de cancelamento',
                       'Situação hoje (frota)', 'Cliente hoje (frota)', 'Ação'],
              linhasA, [20, 10, 19, 26, 11, 55, 28, 12, 20, 26, 55], moedas=(5,),
              titulo=f'Medição de {mes_txt(MES_ATUAL)}: situação de cada rastreador cobrado',
              nota='Uso em agosto/2026 (CEABS: "serviços prestados em agosto de 2026"), NFS-e 10276/10277 de 02/09, vencimento 25/09/2026.')
pinta_status(wsA, 'A', 5, ultA)

# ---------- Contestar CEABS ----------
wsC = wb.create_sheet('Contestar CEABS', 2)
cont = ler('contestar.csv')
linhasC = [[c['mes_uso'], c['tipo'], c['placa'], c['chassi'], c['codigo_ceabs'], num(c['valor']), num(c['valor_contestar']),
            c['motivo'], c['pedido_cancelamento'], c['evidencia']] for c in cont]
for e in eq:
    if e['acao'].startswith('CONTESTAR') or 'não pagar' in e['acao']:
        linhasC.append(['out/26 (aviso de 13/10)', 'Equipamento não devolvido', e['placa'], e['chassi'], e['serie'], num(e['valor']),
                        num(e['valor']), e['acao'], e['pedidos_retirada'], f'Taxa de rescisão cobrada em {e["taxa_rescisao_cobrada"]}; '
                        'planilha "LOCADORA GRILLO E RI - 0130.xlsx" de 13/10/2026'])
ultC = tabela(wsC, 4, ['Mês de uso', 'Tipo', 'Placa', 'Chassi', 'Código CEABS / série', 'Valor cobrado', 'Valor a contestar', 'Motivo',
                       'Pedido de cancelamento', 'Evidência'], linhasC, [12, 18, 10, 19, 16, 12, 12, 55, 14, 60], moedas=(6, 7),
              titulo='Itens para a CEABS conferir (contestação)',
              nota='Lista pronta para anexar ao e-mail de conferência. Abril e agosto/2025 não têm planilha de medição nos e-mails: '
                   'pedir à CEABS também a conferência desses meses para os mesmos veículos.')
r = ultC + 1
wsC.cell(r, 5, 'Total').font = Font(bold=True)
for col in (6, 7):
    c = wsC.cell(r, col, f'=SUBTOTAL(9,{get_column_letter(col)}5:{get_column_letter(col)}{ultC})')
    c.number_format, c.font = BRL, Font(bold=True)

# ---------- Equipamentos ----------
wsE = wb.create_sheet('Equipamentos', 3)
linhasE = [[e['placa'], e['chassi'], e['modelo'], e['dispositivo'], e['serie'], num(e['valor']), e['situacao_frota'], e['data_venda'],
            e['comprador'], e['pedidos_retirada'], e['taxa_rescisao_cobrada'], e['retirada_planilha_monitoramento'], e['acao']] for e in eq]
ultE = tabela(wsE, 4, ['Placa', 'Chassi', 'Modelo', 'Dispositivo', 'Série', 'Valor', 'Situação (frota)', 'Venda', 'Comprador',
                       'Pedidos de retirada', 'Taxa de rescisão cobrada', 'Planilha de retirada (Monitoramento)', 'Ação'],
              linhasE, [10, 19, 24, 12, 13, 9, 14, 11, 18, 22, 12, 22, 70], moedas=(6,),
              titulo='Cobrança de equipamentos não devolvidos: R$ 4.800 a partir de 20/10/2026',
              nota='Aviso da CEABS de 13/10/2026 (lucas.ribeiro@ceabs.com.br): sem pedido de retirada ao corporativo@ceabs.com.br, '
                   'fatura automática a partir do dia 20, boleto 30 dias após a NF.')
wsE.conditional_formatting.add(f'M5:M{ultE}', FormulaRule(formula=['LEFT($M5,9)="CONTESTAR"'], fill=PatternFill('solid', fgColor=COR['CONTESTAR'])))
wsE.conditional_formatting.add(f'M5:M{ultE}', FormulaRule(formula=['LEFT($M5,7)="AGENDAR"'], fill=PatternFill('solid', fgColor=COR['CANCELAR E RETIRAR'])))

# ---------- Retiradas ----------
wsR = wb.create_sheet('Retiradas', 4)
linhasR = [[r['grupo'], r['placa'], r['chassi'], r['modelo'], r['situacao'], r['local'], r['cancelamento'], r['retirada'], r['acao']]
           for r in sorted(ret, key=lambda r: (r['grupo'], r['situacao'].startswith('Retirado'), r['placa']))]
ultR = tabela(wsR, 4, ['Grupo', 'Placa', 'Chassi', 'Modelo', 'Situação', 'Local', 'Cancelamento', 'Retirada', 'Ação'],
              linhasR, [34, 10, 19, 26, 50, 26, 12, 11, 60],
              titulo='Rastreadores: cancelamento e retirada',
              nota='1) sem uso na medição de ago/26: solicitar cancelamento e retirada; 2) já cancelados (taxa de rescisão cobrada desde jun/25): '
                   'situação da retirada pela medição da CEABS e pela planilha do Monitoramento; 3) cancelamento pedido em setembro/26.')
wsR.conditional_formatting.add(f'E5:E{ultR}', FormulaRule(formula=['LEFT($E5,3)="NÃO"'], fill=PatternFill('solid', fgColor=COR['CONTESTAR'])))
wsR.conditional_formatting.add(f'E5:E{ultR}', FormulaRule(formula=['LEFT($E5,8)="Retirado"'], fill=PatternFill('solid', fgColor=COR['OK'])))
wsR.conditional_formatting.add(f'A5:A{ultR}', FormulaRule(formula=['LEFT($A5,1)="1"'], fill=PatternFill('solid', fgColor=COR['CANCELAR E RETIRAR'])))

# ---------- Faturas ----------
wsF = wb.create_sheet('Faturas', 5)
colsF = [('Mês de uso', 'mes_uso', 10), ('Tipo', 'tipo_cobranca', 18), ('Documento', 'documento', 28), ('Emissão', 'emissao', 11),
         ('Vencimento', 'vencimento', 11), ('Vencimento padrão', 'vencimento_padrao', 11), ('Alerta de vencimento', 'alerta_vencimento', 34),
         ('Valor bruto', 'valor_bruto', 12), ('Valor líquido', 'valor_liquido', 12), ('Pago em', 'pago_em', 11),
         ('Situação', 'situacao', 20), ('Dias de atraso', 'dias_atraso', 9), ('Conferência fiscal', 'conferencia_fiscal', 50),
         ('Observação', 'observacao', 50), ('Fonte', 'fonte', 30)]
linhasF = [[num(f[k]) if k in ('valor_bruto', 'valor_liquido', 'dias_atraso') else f[k] for _, k, _ in colsF]
           for f in fat if f['mes_uso'] >= '2024-08']
ultF = tabela(wsF, 4, [c[0] for c in colsF], linhasF, [c[2] for c in colsF], moedas=(8, 9),
              titulo='Faturas CEABS: emissão, vencimento, pagamento e conferência fiscal',
              nota='Padrão de vencimento: dia 3 do 2º mês após o mês de uso. Pagamentos de 2025 pelo controle do Matheus; de 2026 pelo '
                   'CONTROLE CEABS do Monitoramento. Abr/25 e ago/25 sem planilha de medição nos e-mails.')
for sit, cor in (('Em aberto (vencida)', 'F4B6B6'), ('Contestado (em aberto)', 'FFE699'), ('A confirmar', 'D9D9D9'), ('Sem registro de pagamento', 'EDEDED')):
    wsF.conditional_formatting.add(f'K5:K{ultF}', CellIsRule(operator='equal', formula=[f'"{sit}"'], fill=PatternFill('solid', fgColor=cor)))
wsF.conditional_formatting.add(f'G5:G{ultF}', FormulaRule(formula=['LEFT($G5,11)="Vencimento "'], fill=PatternFill('solid', fgColor='FFE699')))

# ---------- Como ler ----------
wsL = wb.create_sheet('Como ler')
wsL.column_dimensions['A'].width = 24
wsL.column_dimensions['B'].width = 110
texto = [
    ('Status', ''),
    ('OK', 'Cobrança procedente. No mês atual, o veículo está em uso: com cliente, reserva SEOP/Casa Civil (passou pelo cliente nos '
           '6 meses anteriores ou volta em até 60 dias), uso interno com abastecimento recente, rastreador ligado/fora das bases no relatório '
           'da CEABS de 17/07/2026, ou em manutenção.'),
    ('CANCELAR E RETIRAR', 'Só no mês atual: rastreador sem uso (vendido, à venda, sem cliente, sinistrado ou uso interno sem uso). '
                           'Solicitar o cancelamento e a retirada do equipamento (aba Retiradas).'),
    ('CONTESTAR', 'Cobrança indevida da CEABS: mês cobrado depois do pedido de cancelamento/retirada, cancelamento já registrado pela CEABS, '
                  'taxa de rescisão acima de R$ 60, visita em duplicidade, equipamento cobrado depois de retirada pedida e taxa paga.'),
    ('VERIFICAR', 'Sem informação suficiente na frota para decidir.'),
    ('', ''),
    ('Mês de uso x vencimento', 'Mês de uso = referência da planilha de medição da CEABS (a CEABS confirmou em 30/09/2026 que a medição de agosto '
                                'é de "serviços prestados em agosto de 2026"). Vencimento = data do boleto da mensalidade daquele mês.'),
    ('Tabelas dinâmicas', 'Três abas "Dinâmica": por mês x status, por cliente/local x status e por veículo x mês. Todas leem a aba Base. '
                          'Use os filtros do topo de cada uma (status, tipo, mês).'),
    ('Fontes', '736 e-mails CEABS do Drive (2022 a out/2026) lidos com anexos; planilhas de medição; NFS-e (Mastersaf e PDFs); '
               'faturas de serviço; controle de pagamentos (Matheus) e CONTROLE CEABS (Monitoramento); base da frota (veículos, passagens, '
               'atendimentos, vendas); diesel; Ticketlog (jun/24 a set/26); relatório de posição da CEABS de 17/07/2026; '
               'planilhas de retirada e de veículos de serviço do Monitoramento.'),
    ('Limitações', 'Sem planilha de medição de abr/25 e ago/25 (só as notas). Jun/26 veio do CONTROLE CEABS do Monitoramento. O e-mail '
                   '"Processo Pagamento - CEABS.msg" (10,9 MB) não abriu pelo conector do Drive.'),
]
for i, (a, b) in enumerate(texto, 1):
    wsL.cell(i, 1, a).font = Font(bold=True, color=AZUL if not b else '000000')
    c = wsL.cell(i, 2, b)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if a in COR:
        wsL.cell(i, 1).fill = PatternFill('solid', fgColor=COR[a])
    wsL.row_dimensions[i].height = 15 if not b else max(15, 15 * (len(b) // 105 + 1))

for w in wb.worksheets:
    w.page_setup.orientation = 'landscape'
    w.page_setup.fitToWidth, w.page_setup.fitToHeight = 1, 0
    w.sheet_properties.pageSetUpPr.fitToPage = True
    w.sheet_properties.tabColor = {'Painel': AZUL, 'Contestar CEABS': 'C00000', 'Equipamentos': 'C00000'}.get(w.title)
wb.save(SAIDA)
print('ok:', SAIDA, '| Base', N - 1, 'linhas | abas:', [w.title for w in wb.worksheets])
